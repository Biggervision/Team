#!/usr/bin/env python3
"""Add a pie + bar chart to the top of every operational tab of a v1
168-Hour Executive OS workbook (the user's uploaded copy), without touching
any cell, formula, validation or formatting.

Both charts read the live LIFE AREA / HOURS table on 📊 Dashboard
(L7:M21 = header + 13 core categories + Free Time), so they update the
moment any input changes.

Usage: python3 add_tab_charts.py <in.xlsx> <out.xlsx>
"""
import sys
import openpyxl
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList

IN, OUT = sys.argv[1], sys.argv[2]
ORANGE = 'F6A21D'

wb = openpyxl.load_workbook(IN)
dash = wb['📊 Dashboard']

# Dashboard life-area table: header row 7, categories 8..20, Free Time 21
DATA = dict(min_col=13, min_row=7, max_row=21)   # M7:M21 (with header)
CATS = dict(min_col=12, min_row=8, max_row=21)   # L8:L21

# sheet -> (pie anchor, bar anchor) in empty space at the top of each page
PLACEMENTS = {
    '📝 Weekly Tracker':  ('I2', 'Q2'),
    '⚙️ Settings':        ('L2', 'T2'),
    '📅 Daily Tracker':   ('N2', 'V2'),
    '📈 Weekly Review':   ('I2', 'Q2'),
    '🗓️ Monthly Tracker': ('I2', 'Q2'),
    '🎯 Annual Tracker':  ('I2', 'Q2'),
}

for name, (pie_at, bar_at) in PLACEMENTS.items():
    if name not in wb.sheetnames:
        print('skip (missing):', name)
        continue
    ws = wb[name]

    pie = PieChart()
    pie.title = 'Time split by %'
    pie.add_data(Reference(dash, **DATA), titles_from_data=True)
    pie.set_categories(Reference(dash, **CATS))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.height, pie.width = 7.2, 9.6
    ws.add_chart(pie, pie_at)

    bar = BarChart()
    bar.type = 'bar'
    bar.title = 'Hours by life area'
    bar.add_data(Reference(dash, **DATA), titles_from_data=True)
    bar.set_categories(Reference(dash, **CATS))
    bar.legend = None
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.numFmt = '0.0'
    bar.height, bar.width = 7.2, 9.6
    s = bar.series[0]
    s.graphicalProperties.solidFill = ORANGE
    ws.add_chart(bar, bar_at)
    print('charts added:', name)

wb.save(OUT)
print('saved', OUT)
