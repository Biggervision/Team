#!/usr/bin/env python3
"""Build ASCENSION OS(tm) v2.0 as a ready-to-import .xlsx.

Output: Ascension-OS-v2.xlsx (repo root).
Import: Google Drive -> New -> File upload -> open with Google Sheets.

v2 features in pure formulas (no scripts):
- The same executive KPI strip on top of EVERY operational sheet
  (used / remaining / weekly / monthly / annual / productive / free /
  utilization / balance) + a live logged today-week-month-year line.
- Tabs: Dashboard, Ideal Week, Current Week, Daily Tracker (date log),
  Monthly Review, Annual Review, Settings, Start Here.
- New life areas: Deep Work, Meetings, Business Development, Content Creation.
LibreOffice-safe formulas only (SUMIFS / INDEX / MATCH / nested IF).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

OUT = '/home/user/Team/Ascension-OS-v2.xlsx'

TEAL, TEAL_DARK = '2BB5AE', '17817B'
BLUE_BOX, INPUT_BG, INPUT_LN = 'A9C7E7', 'DDEBF7', '9DC3E6'
INK, SLATE, FAINT = '1F2937', '5B6B7C', '94A3B8'
ORANGE, GREEN, GREEN_BG = 'F6A21D', '2E9E4F', 'D9F2E3'
RED, RED_BG, AMBER, AMBER_BG = 'C0392B', 'FADBD8', 'B7791F', 'FCEFC7'
LINE, CARD = 'BFC9D4', 'F3F6F9'

DASH, IDL, CW, DT = '📊 Dashboard', '🌟 Ideal Week', '📆 Current Week', '📅 Daily Tracker'
MR, AR, SET, GUI = '🗓️ Monthly Review', '🎯 Annual Review', '⚙️ Settings', '📖 Start Here'

F_BODY  = Font(name='Arial', size=10, color=INK)
F_MUTED = Font(name='Arial', size=10, color=SLATE)
F_BOLD  = Font(name='Arial', size=10, color=INK, bold=True)
F_BAND  = Font(name='Arial', size=11, color='FFFFFF', bold=True)
F_HEAD  = Font(name='Arial', size=9,  color='FFFFFF', bold=True)
F_TITLE = Font(name='Arial', size=16, color=INK, bold=True)
F_SUB   = Font(name='Arial', size=10, color=SLATE, italic=True)
F_KPI_L = Font(name='Arial', size=7,  color=FAINT, bold=True)
F_KPI_V = Font(name='Arial', size=13, color=INK, bold=True)

FILL_TEAL, FILL_INPUT = PatternFill('solid', fgColor=TEAL), PatternFill('solid', fgColor=INPUT_BG)
FILL_BOX, FILL_CARD = PatternFill('solid', fgColor=BLUE_BOX), PatternFill('solid', fgColor=CARD)
thin = Side(style='thin', color=LINE)
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
in_side = Side(style='thin', color=INPUT_LN)
BORDER_IN = Border(top=in_side, bottom=in_side, left=in_side, right=in_side)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left', vertical='center')
RIGHT  = Alignment(horizontal='right', vertical='center')

wb = openpyxl.Workbook()
wb.remove(wb.active)

def sheet(name, tab, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

def band(ws, row, c1, c2, text):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row=row, column=c1, value=text)
    c.font, c.alignment = F_BAND, LEFT
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = FILL_TEAL
    ws.row_dimensions[row].height = 22

def put(ws, row, col, value, font=F_BODY, align=LEFT, fmt=None, fill=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment = font, align
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = border
    return c

def input_cell(ws, row, col, value, fmt='0.##'):
    return put(ws, row, col, value, font=F_BOLD, align=CENTER, fmt=fmt,
               fill=FILL_INPUT, border=BORDER_IN)

def header_cells(ws, row, c1, values):
    for i, h in enumerate(values):
        put(ws, row, c1 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
    ws.cell(row=row, column=c1).alignment = LEFT

# ---------------------------------------------------------------------------
# Fixed layout constants (Settings is the hub every sheet references)
# ---------------------------------------------------------------------------
ST = f"'{SET}'!"
G_WD, G_SG, G_WT = f'{ST}$D$6', f'{ST}$D$7', f'{ST}$D$8'          # goals
C_DAY, C_WEEK = f'{ST}$D$10', f'{ST}$D$11'                        # constants
C_MON, C_YEAR = f'{ST}$D$12', f'{ST}$D$13'
M_FIRST, N_CORE, N_CUSTOM = 17, 17, 5                             # master table
M_LAST = M_FIRST + N_CORE + N_CUSTOM - 1                          # 38
USED, FREE, PROD = f'{ST}$D$40', f'{ST}$D$41', f'{ST}$D$42'
NAMES = f'{ST}$C${M_FIRST}:$C${M_LAST}'
HOURS = f'{ST}$D${M_FIRST}:$D${M_LAST}'
PRODS = f'{ST}$G${M_FIRST}:$G${M_LAST}'

CATS = [  # key, emoji, name, productive, min, max  (rows 17..33)
    ('SLEEP','😴','Sleep',False,49,63), ('WORK','💼','Work',True,'',40),
    ('DEEP','🎯','Deep Work',True,5,''), ('MEET','🗣️','Meetings',True,'',15),
    ('FOOD','🍕','Food',False,3.5,21), ('CHOR','🧹','Household Chores',False,'',15),
    ('CHILD','👶','Childcare',False,'',''), ('FIT','🏋️','Fitness',True,3,20),
    ('ENT','📺','Entertainment',False,'',25), ('REL','❤️','Relationships',False,3,''),
    ('LEARN','📚','Learning',True,1,''), ('PD','🌱','Personal Development',True,1,''),
    ('SPIRIT','🙏','Spiritual',False,'',''), ('TRAVEL','🚗','Travel',False,'',''),
    ('ADMIN','🗂️','Admin',True,'',''), ('BIZ','📈','Business Development',True,'',''),
    ('CONT','🎬','Content Creation',True,'',''),
]
ROW = {c[0]: M_FIRST + i for i, c in enumerate(CATS)}
def SD(key): return f'{ST}$D${ROW[key]}'

# Daily log geometry
L1, LN = 10, 150
LR = f"'{DT}'!"
LOG_DATE = f'{LR}$B${L1}:$B${L1+LN-1}'
LOG_CAT  = f'{LR}$D${L1}:$D${L1+LN-1}'
LOG_ACT  = f'{LR}$F${L1}:$F${L1+LN-1}'
LOG_PROD = f'{LR}$H${L1}:$H${L1+LN-1}'

# Current Week geometry
CW1 = 12
CWL = CW1 + N_CORE + N_CUSTOM - 1                                 # 33
CWR = f"'{CW}'!"
WS_CELL = f'{CWR}$M$9'
CW_EFF = f'{CWR}$E${CW1}:$E${CWL}'
CW_ACT = f'{CWR}$F${CW1}:$F${CWL}'
CW_VAR = f'{CWR}$G${CW1}:$G${CWL}'

WEEK_ACT = f'SUMIFS({LOG_ACT},{LOG_DATE},">="&{WS_CELL},{LOG_DATE},"<"&({WS_CELL}+7))'

BALANCE = (f'=ROUND(100*(({SD("SLEEP")}>={G_SG}*7*0.95)*({SD("SLEEP")}<=63)'
           f'+(({SD("WORK")}+{SD("DEEP")}+{SD("MEET")})<=55)*1'
           f'+({SD("FIT")}>=3)*1+({SD("REL")}>=3)*1'
           f'+(({SD("LEARN")}+{SD("PD")})>=2)*1'
           f'+({FREE}>=5)*1+({FREE}>=0)*1)/7,0)')

def kpi_strip(ws, title, subtitle):
    """Rows 2-6: title, subtitle, KPI cards, live logged line. Returns next row."""
    put(ws, 2, 2, title, F_TITLE)
    put(ws, 3, 2, subtitle, F_SUB)
    kpis = [
        ('HOURS USED',  f'={USED}', '0.0'),
        ('REMAINING',   f'=MAX(0,{C_WEEK}-{WEEK_ACT})', '0.0'),
        ('WEEKLY HRS',  f'={C_WEEK}', '0'),
        ('MONTHLY HRS', f'={USED}*{C_MON}', '#,##0'),
        ('ANNUAL HRS',  f'={USED}*{C_YEAR}', '#,##0'),
        ('PRODUCTIVE',  f'={PROD}', '0.0'),
        ('FREE TIME',   f'={FREE}', '0.0'),
        ('UTILIZATION', f'={USED}/{C_WEEK}', '0.0%'),
        ('BALANCE',     BALANCE, '0'),
    ]
    for i, (lab, formula, fmt) in enumerate(kpis):
        c = 2 + i
        put(ws, 4, c, lab, F_KPI_L, CENTER, fill=FILL_CARD, border=BORDER)
        put(ws, 5, c, formula, F_KPI_V, CENTER, fmt, fill=FILL_CARD, border=BORDER)
    ws.row_dimensions[4].height = 13
    ws.row_dimensions[5].height = 22
    ws.merge_cells(start_row=6, start_column=2, end_row=6, end_column=10)
    put(ws, 6, 2,
        f'="⏱  Logged — today: "&TEXT(SUMIFS({LOG_ACT},{LOG_DATE},TODAY()),"0.0")'
        f'&" h  ·  this week: "&TEXT({WEEK_ACT},"0.0")'
        f'&" h  ·  this month: "&TEXT(SUMIFS({LOG_ACT},{LOG_DATE},">="&EOMONTH(TODAY(),-1)+1,{LOG_DATE},"<="&EOMONTH(TODAY(),0)),"0.0")'
        f'&" h  ·  this year: "&TEXT(SUMIFS({LOG_ACT},{LOG_DATE},">="&DATE(YEAR(TODAY()),1,1)),"0.0")&" h"',
        Font(name='Arial', size=9, color=SLATE))
    bal = ws.cell(row=5, column=10)
    ws.conditional_formatting.add(bal.coordinate, FormulaRule(
        formula=[f'{bal.coordinate}>=80'], font=Font(name='Arial', size=13, bold=True, color=GREEN)))
    ws.conditional_formatting.add(bal.coordinate, FormulaRule(
        formula=[f'{bal.coordinate}<60'], font=Font(name='Arial', size=13, bold=True, color=RED)))
    return 8

def cat_row_formulas(ws, r, i, extra):
    """Category + ideal columns shared by CW / Monthly / Annual tables."""
    sr = M_FIRST + i
    put(ws, r, 2, f'=IF({ST}C{sr}="","",{ST}B{sr}&"  "&{ST}C{sr})', F_BODY, border=BORDER)
    put(ws, r, 3, f'=IF($B{r}="","",{ST}D{sr})', F_BODY, CENTER, '0.0', border=BORDER)
    extra(r, sr)

# ===========================================================================
# ⚙️ SETTINGS
# ===========================================================================
st = sheet(SET, TEAL_DARK, [2, 5, 26, 11, 10, 9, 11, 9, 9, 10])
put(st, 2, 2, '⚙️ Settings — goals & master data', F_TITLE)
put(st, 3, 2, 'Light-blue cells are yours. Custom categories flow into every sheet automatically.', F_SUB)
band(st, 5, 2, 10, '  🎯  GOALS & SETUP')
goals = [('Work days per week', 6, 'Drives lunch, commute, deep-work formulas.'),
         ('Sleep goal (hours per night)', 7.5, 'Feeds the Balance Score & sleep warnings.'),
         ('Weekly productive target (h)', 45, 'Goal-progress measures against this.')]
for i, (lab, val, note) in enumerate(goals):
    r = 6 + i
    put(st, r, 2, lab, F_MUTED); st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    input_cell(st, r, 4, val)
    put(st, r, 5, note, Font(name='Arial', size=8, color=FAINT))
    st.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
band(st, 9, 2, 10, '  ⏱️  CONSTANTS')
consts = [('Hours per day', 24), ('Hours per week', '=D10*7'),
          ('Weeks per month', 4), ('Weeks per year', 52)]
for i, (lab, val) in enumerate(consts):
    r = 10 + i
    put(st, r, 2, lab, F_MUTED); st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = put(st, r, 4, val, F_BOLD, CENTER, '0.###')
st['D12'].comment = Comment('Executive convention (4 wk/month, 52 wk/year). Use 4.345 / 52.18 for calendar-accurate maths.', 'Ascension OS')

band(st, 15, 2, 10, '  🗂️  MASTER CATEGORY TABLE')
header_cells(st, 16, 2, ['', 'Category', 'Ideal h/wk', '% of week', 'Per day',
                         'Productive', 'Min h/wk', 'Max h/wk', 'Status'])
plan_cell = {}   # key -> Ideal Week total ref, filled while building Ideal Week
for i, (key, emoji, name, prod, mn, mx) in enumerate(CATS):
    r = M_FIRST + i
    put(st, r, 2, emoji, F_BODY, CENTER, border=BORDER)
    put(st, r, 3, name, F_BODY, border=BORDER)
    put(st, r, 4, None, F_BOLD, CENTER, '0.0', border=BORDER)   # formula set later
    put(st, r, 7, prod, F_BODY, CENTER, border=BORDER)
    put(st, r, 8, mn if mn != '' else None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
    put(st, r, 9, mx if mx != '' else None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
for i in range(N_CUSTOM):
    r = M_FIRST + N_CORE + i
    input_cell(st, r, 2, None); st.cell(row=r, column=2).alignment = CENTER
    input_cell(st, r, 3, None); st.cell(row=r, column=3).alignment = LEFT
    input_cell(st, r, 4, None, '0.0')
    put(st, r, 7, False, F_BODY, CENTER, fill=FILL_INPUT, border=BORDER_IN)
    put(st, r, 8, None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
    put(st, r, 9, None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
st.cell(row=M_FIRST + N_CORE, column=3).comment = Comment(
    'Type a custom category name + weekly hours — it appears on every sheet.', 'Ascension OS')
for r in range(M_FIRST, M_LAST + 1):
    put(st, r, 5, f'=IF($C{r}="","",$D{r}/{C_WEEK})', F_BODY, CENTER, '0.0%', border=BORDER)
    put(st, r, 6, f'=IF($C{r}="","",$D{r}/7)', F_BODY, CENTER, '0.0', border=BORDER)
    put(st, r, 10, f'=IF($C{r}="","",IF(AND($H{r}<>"",$D{r}<$H{r}),"🔻 Low",'
                   f'IF(AND($I{r}<>"",$D{r}>$I{r}),"🔺 High","✅ OK")))',
        F_BODY, CENTER, border=BORDER)
put(st, 40, 3, 'Total allocated', F_BOLD, RIGHT)
put(st, 40, 4, f'=SUM(D{M_FIRST}:D{M_LAST})', F_BOLD, CENTER, '0.0')
put(st, 41, 3, 'Free time (168 − allocated)', F_BOLD, RIGHT)
put(st, 41, 4, f'={C_WEEK}-D40', F_BOLD, CENTER, '0.0')
put(st, 42, 3, 'Productive hours', F_BOLD, RIGHT)
put(st, 42, 4, f'=SUMPRODUCT(--(G{M_FIRST}:G{M_LAST}=TRUE),D{M_FIRST}:D{M_LAST})', F_BOLD, CENTER, '0.0')
st.conditional_formatting.add(f'D{M_FIRST}:D{M_LAST}', FormulaRule(
    formula=[f'AND($I{M_FIRST}<>"",$D{M_FIRST}>$I{M_FIRST})'],
    fill=PatternFill('solid', fgColor=RED_BG), font=Font(name='Arial', size=10, color=RED, bold=True)))
st.conditional_formatting.add(f'D{M_FIRST}:D{M_LAST}', FormulaRule(
    formula=[f'AND($H{M_FIRST}<>"",$D{M_FIRST}<$H{M_FIRST})'],
    fill=PatternFill('solid', fgColor=AMBER_BG), font=Font(name='Arial', size=10, color=AMBER, bold=True)))
st.conditional_formatting.add('D41', FormulaRule(
    formula=['$D$41<0'], fill=PatternFill('solid', fgColor=RED_BG),
    font=Font(name='Arial', size=10, color=RED, bold=True)))
st.freeze_panes = 'A17'

wb.defined_names['CatList'] = DefinedName('CatList',
    attr_text=f"'{SET}'!$C${M_FIRST}:$C${M_LAST}")

# ===========================================================================
# 🌟 IDEAL WEEK — guided sentences
# ===========================================================================
idl = sheet(IDL, TEAL, [2, 9, 9, 9, 9, 9, 9, 9, 9, 9])
for c, w in zip(range(2, 8), [30, 9, 40, 21, 9, 46]):
    idl.column_dimensions[get_column_letter(c)].width = w
row = kpi_strip(idl, '🌟 Ideal Week — design your 168 hours',
    'Only edit the light-blue cells. Every sheet in the OS measures reality against this design.')
r = row + 1

def line(label, val, unit, calc=None, tail=None, ref_display=None):
    global r
    put(idl, r, 2, label)
    if ref_display:
        put(idl, r, 3, ref_display, F_MUTED, CENTER, '0')
    else:
        input_cell(idl, r, 3, val)
    put(idl, r, 4, unit, F_MUTED)
    if calc:
        put(idl, r, 5, 'This means I spend', F_MUTED, RIGHT)
        put(idl, r, 6, calc(r), F_BOLD, CENTER, '0.0')
        put(idl, r, 7, tail, F_MUTED)
    r += 1
    return r - 1

prev_leaves = None
def close(key, emoji, name, f_rows, tail):
    global r, prev_leaves
    put(idl, r, 5, f'{emoji} In total, I spend', F_BOLD, RIGHT)
    put(idl, r, 6, '=' + '+'.join(f'F{x}' for x in f_rows), F_BOLD, CENTER, '0.0')
    put(idl, r, 7, f'="{tail}  ≈ "&TEXT($F${r}*{C_MON},"0")&" h/month · "'
                   f'&TEXT($F${r}*{C_YEAR},"0")&" h/year · "&TEXT($F${r}/{C_WEEK},"0.0%")&" of your week"',
        Font(name='Arial', size=9, color=SLATE))
    plan_cell[key] = f"'{IDL}'!$F${r}"
    tr = r; r += 1
    put(idl, r, 5, '✅ That leaves me with', F_BOLD, RIGHT)
    base = f'{C_WEEK}' if prev_leaves is None else f'F{prev_leaves}'
    lv = put(idl, r, 6, f'={base}-F{tr}', F_BOLD, CENTER, '0.0')
    lv.font = Font(name='Arial', size=10, bold=True, color=GREEN)
    put(idl, r, 7, 'hours per week for everything else.', F_MUTED)
    prev_leaves = r; r += 2

def open_sec(emoji, name):
    global r
    band(idl, r, 2, 7, f'  {emoji}  {name}'); r += 1

open_sec('😴', 'SLEEP')
a = line('I sleep for around', 8, 'hours per night.', lambda rr: f'=C{rr}*7', 'hours per week sleeping.')
b = line('I wind down for', 10, 'minutes before sleeping.', lambda rr: f'=C{rr}*7/60', 'hours per week preparing to sleep.')
c0 = line('My wake-up routine takes', 0, 'minutes each morning.', lambda rr: f'=C{rr}*7/60', 'hours per week waking up.')
close('SLEEP', '😴', 'Sleep', [a, b, c0], 'hours per week on sleep-related things.')

open_sec('💼', 'WORK')
line('I work', None, 'days per week  (change in ⚙️ Settings).', ref_display=f'={G_WD}')
w1 = line('Outside meetings & deep work, general work takes me', 3, 'hours per working day.',
          lambda rr: f'=C{rr}*{G_WD}', 'hours per week on general work.')
w2 = line('My lunch break is', 30, 'minutes per working day.',
          lambda rr: f'=C{rr}*{G_WD}/60', 'hours per week on my lunch break.')
w3 = line('My commute is', 0, 'minutes each way.',
          lambda rr: f'=C{rr}*2*{G_WD}/60', 'hours per week commuting.')
w4 = line('Getting ready for work takes', 10, 'minutes per working day.',
          lambda rr: f'=C{rr}*{G_WD}/60', 'hours per week getting ready.')
w5 = line('Switching off after work takes', 10, 'minutes per working day.',
          lambda rr: f'=C{rr}*{G_WD}/60', 'hours per week switching off.')
close('WORK', '💼', 'Work', [w1, w2, w3, w4, w5], 'hours per week working.')

open_sec('🎯', 'DEEP WORK')
d1 = line('I do', 2, 'hours of deep, focused work per working day.',
          lambda rr: f'=C{rr}*{G_WD}', 'hours per week in deep work.')
close('DEEP', '🎯', 'Deep Work', [d1], 'hours per week of deep work.')

open_sec('🗣️', 'MEETINGS')
m1 = line('I spend', 5, 'hours in meetings per week.', lambda rr: f'=C{rr}', 'hours per week in meetings.')
close('MEET', '🗣️', 'Meetings', [m1], 'hours per week in meetings.')

open_sec('🍕', 'FOOD')
f1 = line('Breakfast takes', 15, 'minutes per day (eating + prep).', lambda rr: f'=C{rr}*7/60', 'hours per week on breakfast.')
f2 = line('Lunch takes', 40, 'minutes per non-working day.', lambda rr: f'=C{rr}*(7-{G_WD})/60', 'hours per week on lunch (non-work days).')
f3 = line('Dinner takes', 30, 'minutes per day (eating + prep).', lambda rr: f'=C{rr}*7/60', 'hours per week on dinner.')
close('FOOD', '🍕', 'Food', [f1, f2, f3], 'hours per week on food.')

open_sec('🧹', 'HOUSEHOLD CHORES')
g1 = line('I do groceries', 2, 'times per week.')
g2 = line('Each grocery run takes', 60, 'minutes (incl. travel).', lambda rr: f'=C{g1}*C{rr}/60', 'hours per week on groceries.')
g3 = line('I clean the house', 1, 'times per week.')
g4 = line('Each cleaning session takes', 60, 'minutes.', lambda rr: f'=C{g3}*C{rr}/60', 'hours per week on cleaning.')
g5 = line('Other misc. chores take', 1, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on misc. chores.')
close('CHOR', '🧹', 'Household Chores', [g2, g4, g5], 'hours per week on chores.')

open_sec('👶', 'CHILDCARE')
k1 = line('Each weekday I spend', 0, 'hours actively on childcare.', lambda rr: f'=C{rr}*5', 'hours on weekdays on childcare.')
k2 = line('Each weekend day I spend', 0, 'hours actively on childcare.', lambda rr: f'=C{rr}*2', 'hours on weekends on childcare.')
k3 = line('Child admin & logistics take', 0, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on child admin.')
close('CHILD', '👶', 'Childcare', [k1, k2, k3], 'hours per week on childcare.')

open_sec('🏋️', 'FITNESS')
t1 = line('Each week, I spend', 8, 'hours exercising (incl. commute time).', lambda rr: f'=C{rr}', 'hours per week exercising.')
t2 = line('Each week, I spend', 2, 'hours on other fitness-related things.', lambda rr: f'=C{rr}', 'hours per week on other fitness.')
close('FIT', '🏋️', 'Fitness', [t1, t2], 'hours per week on fitness.')

open_sec('📺', 'ENTERTAINMENT')
e1 = line('Each week, I spend', 0, 'hours watching TV shows / movies.', lambda rr: f'=C{rr}', 'hours per week on TV / movies.')
e2 = line('Each week, I spend', 7, 'hours on social media apps.', lambda rr: f'=C{rr}', 'hours per week on social media.')
e3 = line('Each week, I spend', 2, 'hours reading.', lambda rr: f'=C{rr}', 'hours per week reading.')
e4 = line('Each week, I spend', 3, 'hours on other entertainment.', lambda rr: f'=C{rr}', 'hours per week on other entertainment.')
close('ENT', '📺', 'Entertainment', [e1, e2, e3, e4], 'hours per week on entertainment.')

open_sec('❤️', 'RELATIONSHIPS')
r1 = line('Each week, I spend', 3, 'hours on quality family time.', lambda rr: f'=C{rr}', 'hours per week on family time.')
r2 = line('Each week, I spend', 3, 'hours on other socialising time.', lambda rr: f'=C{rr}', 'hours per week socialising.')
close('REL', '❤️', 'Relationships', [r1, r2], 'hours per week on relationships.')

open_sec('📚', 'LEARNING')
q1 = line('Courses & structured learning take', 1, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on courses.')
q2 = line('Books, articles & study take', 1, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week studying.')
q3 = line('I listen to podcasts / audiobooks for', 20, 'minutes per day.', lambda rr: f'=C{rr}*7/60', 'hours per week on podcasts.')
close('LEARN', '📚', 'Learning', [q1, q2, q3], 'hours per week on learning.')

open_sec('🌱', 'PERSONAL DEVELOPMENT')
p1 = line('I journal for', 10, 'minutes per day.', lambda rr: f'=C{rr}*7/60', 'hours per week journaling.')
p2 = line('I meditate for', 10, 'minutes per day.', lambda rr: f'=C{rr}*7/60', 'hours per week meditating.')
p3 = line('Weekly planning & review takes', 1, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week planning.')
close('PD', '🌱', 'Personal Development', [p1, p2, p3], 'hours per week on personal development.')

open_sec('🙏', 'SPIRITUAL')
s1 = line('Daily practice takes', 0, 'minutes per day.', lambda rr: f'=C{rr}*7/60', 'hours per week on daily practice.')
s2 = line('Services & community take', 0, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on services.')
close('SPIRIT', '🙏', 'Spiritual', [s1, s2], 'hours per week on spiritual life.')

open_sec('🚗', 'TRAVEL')
v1 = line('Errands & getting around (non-commute) take', 1.5, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on errands.')
v2 = line('Trips & other travel take', 0, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on trips.')
close('TRAVEL', '🚗', 'Travel', [v1, v2], 'hours per week travelling.')

open_sec('🗂️', 'ADMIN')
a1 = line('Finances, email & life admin take', 2, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week on life admin.')
a2 = line('Appointments & paperwork take', 30, 'minutes per week.', lambda rr: f'=C{rr}/60', 'hours per week on paperwork.')
close('ADMIN', '🗂️', 'Admin', [a1, a2], 'hours per week on admin.')

open_sec('📈', 'BUSINESS DEVELOPMENT')
b1 = line('I spend', 4, 'hours per week on sales, partnerships & growth.', lambda rr: f'=C{rr}', 'hours per week on business development.')
close('BIZ', '📈', 'Business Development', [b1], 'hours per week on business development.')

open_sec('🎬', 'CONTENT CREATION')
c1 = line('I create content for', 6, 'hours per week.', lambda rr: f'=C{rr}', 'hours per week creating content.')
close('CONT', '🎬', 'Content Creation', [c1], 'hours per week creating content.')

idl.freeze_panes = 'A8'
dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='10000',
                    allow_blank=True, showErrorMessage=True, error='Enter a number (0 or more).')
idl.add_data_validation(dv); dv.add(f'C9:C{r}')

# now wire Settings master hours to the Ideal Week totals
for key, cell in plan_cell.items():
    st.cell(row=ROW[key], column=4).value = f'={cell}'

# ===========================================================================
# 📆 CURRENT WEEK
# ===========================================================================
cwx = sheet(CW, TEAL, [2, 26, 9, 9, 9, 9, 9, 10, 9, 9, 2, 2, 12])
row = kpi_strip(cwx, '📆 Current Week — plan, then perform',
    'Leave "My plan" blank to use your Ideal Week. Actuals arrive from the 📅 Daily Tracker automatically.')
cwx.merge_cells('B9:F9')
put(cwx, 9, 2, f'="Week of "&TEXT({WS_CELL},"ddd d mmm")&" – "&TEXT({WS_CELL}+6,"ddd d mmm yyyy")', F_BOLD)
put(cwx, 9, 12, 'starts:', Font(name='Arial', size=8, color=FAINT), RIGHT)
put(cwx, 9, 13, '=TODAY()-WEEKDAY(TODAY(),2)+1', Font(name='Arial', size=8, color=FAINT), CENTER, 'yyyy-mm-dd')
header_cells(cwx, 11, 2, ['Category', 'Ideal', 'My plan', 'Planned', 'Actual',
                          'Variance', 'Remaining', '% of plan'])
for i in range(N_CORE + N_CUSTOM):
    rr = CW1 + i
    def extra(rr, sr):
        input_cell(cwx, rr, 4, None, '0.0')
        put(cwx, rr, 5, f'=IF($B{rr}="","",IF($D{rr}="",$C{rr},$D{rr}))', F_BOLD, CENTER, '0.0', border=BORDER)
        put(cwx, rr, 6, f'=IF($B{rr}="","",SUMIFS({LOG_ACT},{LOG_CAT},{ST}C{sr},{LOG_DATE},">="&{WS_CELL},{LOG_DATE},"<"&({WS_CELL}+7)))',
            F_BOLD, CENTER, '0.0', border=BORDER)
        put(cwx, rr, 7, f'=IF($B{rr}="","",$F{rr}-$E{rr})', F_BODY, CENTER, '0.0', border=BORDER)
        put(cwx, rr, 8, f'=IF($B{rr}="","",MAX(0,$E{rr}-$F{rr}))', F_BODY, CENTER, '0.0', border=BORDER)
        put(cwx, rr, 9, f'=IF(OR($B{rr}="",$E{rr}=0),"",$F{rr}/$E{rr})', F_BODY, CENTER, '0%', border=BORDER)
    cat_row_formulas(cwx, rr, i, extra)
put(cwx, CWL + 2, 2, 'Total planned / logged', F_BOLD, RIGHT)
put(cwx, CWL + 2, 5, f'=SUM(E{CW1}:E{CWL})', F_BOLD, CENTER, '0.0')
put(cwx, CWL + 2, 6, f'=SUM(F{CW1}:F{CWL})', F_BOLD, CENTER, '0.0')
put(cwx, CWL + 3, 2, 'Free / remaining hours', F_BOLD, RIGHT)
put(cwx, CWL + 3, 5, f'={C_WEEK}-E{CWL+2}', F_BOLD, CENTER, '0.0')
put(cwx, CWL + 3, 6, f'={C_WEEK}-F{CWL+2}', F_BOLD, CENTER, '0.0')
put(cwx, CWL + 4, 2, 'Utilization', F_BOLD, RIGHT)
put(cwx, CWL + 4, 5, f'=E{CWL+2}/{C_WEEK}', F_BOLD, CENTER, '0.0%')
put(cwx, CWL + 4, 6, f'=F{CWL+2}/{C_WEEK}', F_BOLD, CENTER, '0.0%')
cwx.conditional_formatting.add(f'F{CW1}:F{CWL}',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=60,
                color=ORANGE, showValue=True))
cwx.conditional_formatting.add(f'I{CW1}:I{CWL}', FormulaRule(
    formula=[f'AND($B{CW1}<>"",$I{CW1}<>"",$I{CW1}>1.2)'],
    fill=PatternFill('solid', fgColor=RED_BG), font=Font(name='Arial', size=10, color=RED)))
cwx.conditional_formatting.add(f'I{CW1}:I{CWL}', FormulaRule(
    formula=[f'AND($B{CW1}<>"",$I{CW1}<>"",$I{CW1}<0.5)'],
    fill=PatternFill('solid', fgColor=AMBER_BG), font=Font(name='Arial', size=10, color=AMBER)))
dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='168',
                    allow_blank=True, showErrorMessage=True, error='Hours must be 0–168.')
cwx.add_data_validation(dv); dv.add(f'D{CW1}:D{CWL}')
cwx.freeze_panes = 'A12'

# ===========================================================================
# 📅 DAILY TRACKER — the execution engine
# ===========================================================================
dt = sheet(DT, TEAL, [2, 12, 6, 24, 9, 9, 34, 8, 2, 15, 8, 8, 8])
row = kpi_strip(dt, '📅 Daily Tracker — the execution engine',
    'Log each block: date, category, actual hours, notes. Week, month and year totals build themselves.')
band(dt, 8, 2, 8, '  📅  DAILY LOG')
band(dt, 8, 10, 13, '  ⚡ TODAY')
header_cells(dt, 9, 2, ['Date', 'Day', 'Category', 'Plan/day', 'Actual h', 'Notes', 'Prod.'])
TODAY_TOTAL = f'SUMIFS({LOG_ACT},{LOG_DATE},TODAY())'
today_rows = [
    ('Daily total', f'={TODAY_TOTAL}', '0.0'),
    ('Remaining today', f'={C_DAY}-{TODAY_TOTAL}', '0.0'),
    ('Productivity score', f'=IFERROR(SUMIFS({LOG_ACT},{LOG_DATE},TODAY(),{LOG_PROD},TRUE)/{TODAY_TOTAL},0)', '0%'),
]
for i, (lab, formula, fmt) in enumerate(today_rows):
    put(dt, 9 + i, 10, lab, F_MUTED, border=BORDER)
    put(dt, 9 + i, 11, formula, F_BOLD, CENTER, fmt, border=BORDER)
band(dt, 13, 10, 13, '  📊  LAST 7 DAYS')
header_cells(dt, 14, 10, ['Day', 'Logged', 'Left', 'Prod %'])
for i in range(7):
    rr = 15 + i
    put(dt, rr, 10, f'=TEXT(TODAY()-{i},"ddd d mmm")', F_MUTED, border=BORDER)
    put(dt, rr, 11, f'=SUMIFS({LOG_ACT},{LOG_DATE},TODAY()-{i})', F_BODY, CENTER, '0.0', border=BORDER)
    put(dt, rr, 12, f'={C_DAY}-K{rr}', Font(name='Arial', size=10, color=FAINT), CENTER, '0.0', border=BORDER)
    put(dt, rr, 13, f'=IFERROR(SUMIFS({LOG_ACT},{LOG_DATE},TODAY()-{i},{LOG_PROD},TRUE)/K{rr},"")', F_BODY, CENTER, '0%', border=BORDER)
dt.conditional_formatting.add('K15:K21', FormulaRule(
    formula=['K15>24'], fill=PatternFill('solid', fgColor=RED_BG),
    font=Font(name='Arial', size=10, color=RED, bold=True)))
dt.conditional_formatting.add('K15:K21', FormulaRule(
    formula=['AND(K15>18,K15<=24)'], fill=PatternFill('solid', fgColor=AMBER_BG),
    font=Font(name='Arial', size=10, color=AMBER, bold=True)))

for i in range(LN):
    rr = L1 + i
    input_cell(dt, rr, 2, None, 'yyyy-mm-dd')
    put(dt, rr, 3, f'=IF($B{rr}="","",TEXT($B{rr},"ddd"))',
        Font(name='Arial', size=10, color=FAINT), CENTER, border=BORDER)
    input_cell(dt, rr, 4, None); dt.cell(row=rr, column=4).alignment = LEFT
    put(dt, rr, 5, f'=IF($D{rr}="","",IFERROR(INDEX({CW_EFF},MATCH($D{rr},{NAMES},0))/7,""))',
        Font(name='Arial', size=10, color=FAINT), CENTER, '0.00', border=BORDER)
    input_cell(dt, rr, 6, None, '0.0#')
    input_cell(dt, rr, 7, None); dt.cell(row=rr, column=7).alignment = LEFT
    put(dt, rr, 8, f'=IF($D{rr}="","",IFERROR(INDEX({PRODS},MATCH($D{rr},{NAMES},0)),FALSE))',
        Font(name='Arial', size=8, color=FAINT), CENTER, border=BORDER)
    for cc in (2, 4, 6, 7):
        dt.cell(row=rr, column=cc).font = Font(name='Arial', size=10, color=INK)
dvd = DataValidation(type='date', allow_blank=True, showErrorMessage=True, error='Enter a date.')
dt.add_data_validation(dvd); dvd.add(f'B{L1}:B{L1+LN-1}')
dvc = DataValidation(type='list', formula1='CatList', allow_blank=True,
                     showErrorMessage=True, error='Pick a category from the list.')
dt.add_data_validation(dvc); dvc.add(f'D{L1}:D{L1+LN-1}')
dvh = DataValidation(type='decimal', operator='between', formula1='0', formula2='24',
                     allow_blank=True, showErrorMessage=True, error='Hours must be 0–24.')
dt.add_data_validation(dvh); dvh.add(f'F{L1}:F{L1+LN-1}')
dt.freeze_panes = 'A10'

# ===========================================================================
# 🗓️ MONTHLY REVIEW
# ===========================================================================
MS, ME = 'EOMONTH(TODAY(),-1)+1', 'EOMONTH(TODAY(),0)'
mr = sheet(MR, TEAL, [2, 26, 9, 10, 11, 9, 10, 7, 2, 8, 8, 8, 8, 8])
row = kpi_strip(mr, '🗓️ Monthly Review',
    'Everything derives automatically from the 📅 Daily Tracker log.')
mr.merge_cells('B9:E9')
put(mr, 9, 2, '="Reviewing:  "&TEXT(TODAY(),"mmmm yyyy")', F_BOLD)
band(mr, 10, 2, 8, '  🗓️  MONTHLY SUMMARY')
band(mr, 10, 10, 14, '  📆  WEEKLY BREAKDOWN')
header_cells(mr, 11, 2, ['Category', 'Ideal h/wk', 'Ideal month', 'Actual month', 'Variance', '% of month', 'Rank'])
header_cells(mr, 11, 10, ['W1', 'W2', 'W3', 'W4', 'W5'])
MR1 = 12
MRL = MR1 + N_CORE + N_CUSTOM - 1
for i in range(N_CORE + N_CUSTOM):
    rr = MR1 + i
    def extra(rr, sr):
        put(mr, rr, 4, f'=IF($B{rr}="","",$C{rr}*{C_MON})', F_BODY, CENTER, '0.0', border=BORDER)
        put(mr, rr, 5, f'=IF($B{rr}="","",SUMIFS({LOG_ACT},{LOG_CAT},{ST}C{sr},{LOG_DATE},">="&{MS},{LOG_DATE},"<="&{ME}))',
            F_BOLD, CENTER, '0.0', border=BORDER)
        put(mr, rr, 6, f'=IF($B{rr}="","",$E{rr}-$D{rr})', F_BODY, CENTER, '0.0', border=BORDER)
        put(mr, rr, 7, f'=IF($B{rr}="","",$D{rr}/({C_WEEK}*{C_MON}))', F_BODY, CENTER, '0.0%', border=BORDER)
        put(mr, rr, 8, f'=IF(OR($B{rr}="",$E{rr}=0),"",RANK($E{rr},$E${MR1}:$E${MRL}))', F_BODY, CENTER, border=BORDER)
        for w in range(5):
            put(mr, rr, 10 + w,
                f'=IF($B{rr}="","",SUMIFS({LOG_ACT},{LOG_CAT},{ST}C{sr},{LOG_DATE},">="&{MS}+{w*7},{LOG_DATE},"<"&MIN({ME}+1,{MS}+{(w+1)*7})))',
                F_BODY, CENTER, '0.0', border=BORDER)
    cat_row_formulas(mr, rr, i, extra)
put(mr, MRL + 2, 2, 'Total', F_BOLD, RIGHT)
put(mr, MRL + 2, 4, f'=SUM(D{MR1}:D{MRL})', F_BOLD, CENTER, '0.0')
put(mr, MRL + 2, 5, f'=SUM(E{MR1}:E{MRL})', F_BOLD, CENTER, '0.0')
band(mr, MRL + 4, 2, 8, '  💡  MONTHLY INSIGHTS')
E_, F_, B_ = f'$E${MR1}:$E${MRL}', f'$F${MR1}:$F${MRL}', f'$B${MR1}:$B${MRL}'
insights = [
    f'=IF(SUM({E_})=0,"Log days on the 📅 Daily Tracker to unlock monthly insights.",'
    f'"🏆  Biggest time investment this month: "&INDEX({B_},MATCH(MAX({E_}),{E_},0))&"  ("&TEXT(MAX({E_}),"0.0")&" h logged)")',
    f'=IF(SUM({E_})=0,"",IF(MAX({F_})<=0,"✅  No category is over its monthly plan.",'
    f'"🔺  Most over plan: "&INDEX({B_},MATCH(MAX({F_}),{F_},0))&"  (+"&TEXT(MAX({F_}),"0.0")&" h)"))',
    f'=IF(SUM({E_})=0,"","🔻  Most behind plan: "&INDEX({B_},MATCH(MIN({F_}),{F_},0))&"  ("&TEXT(MIN({F_}),"0.0")&" h)")',
    f'="⚖️  Month utilization so far: "&TEXT(SUM({E_})/({C_WEEK}*{C_MON}),"0.0%")&" of "&ROUND({C_WEEK}*{C_MON},0)&" hours."',
]
for i, f in enumerate(insights):
    mr.merge_cells(start_row=MRL + 5 + i, start_column=2, end_row=MRL + 5 + i, end_column=8)
    put(mr, MRL + 5 + i, 2, f, F_MUTED)
mr.freeze_panes = 'A12'

# ===========================================================================
# 🎯 ANNUAL REVIEW
# ===========================================================================
YS = 'DATE(YEAR(TODAY()),1,1)'
WKS = f'MAX(1,ROUNDUP((TODAY()-{YS}+1)/7,0))'
ar = sheet(AR, TEAL, [2, 26, 9, 11, 11, 9, 11, 10, 10])
row = kpi_strip(ar, '🎯 Annual Review',
    'Ideal Week × 52 versus what you actually logged. Small weekly numbers become life-sized here.')
band(ar, 8, 2, 9, '  🎯  YEAR — PLAN VS REALITY')
header_cells(ar, 9, 2, ['Category', 'Weekly h', 'Annual plan', 'Actual YTD', 'On pace', 'Annual days', '% of year', 'Balance'])
AR1 = 10
ARL = AR1 + N_CORE + N_CUSTOM - 1
for i in range(N_CORE + N_CUSTOM):
    rr = AR1 + i
    def extra(rr, sr):
        put(ar, rr, 4, f'=IF($B{rr}="","",$C{rr}*{C_YEAR})', F_BODY, CENTER, '#,##0', border=BORDER)
        put(ar, rr, 5, f'=IF($B{rr}="","",SUMIFS({LOG_ACT},{LOG_CAT},{ST}C{sr},{LOG_DATE},">="&{YS}))',
            F_BOLD, CENTER, '#,##0', border=BORDER)
        put(ar, rr, 6, f'=IF(OR($B{rr}="",$C{rr}=0),"",$E{rr}/($C{rr}*{WKS}))', F_BODY, CENTER, '0%', border=BORDER)
        put(ar, rr, 7, f'=IF($B{rr}="","",$D{rr}/{C_DAY})', F_BODY, CENTER, '0.0', border=BORDER)
        put(ar, rr, 8, f'=IF($B{rr}="","",$C{rr}/{C_WEEK})', F_BODY, CENTER, '0.0%', border=BORDER)
        put(ar, rr, 9, f'=IF($B{rr}="","",{ST}J{sr})', F_BODY, CENTER, border=BORDER)
    cat_row_formulas(ar, rr, i, extra)
band(ar, ARL + 2, 2, 9, '  📖  YEAR IN REVIEW')
E_, B_ = f'$E${AR1}:$E${ARL}', f'$B${AR1}:$B${ARL}'
year_lines = [
    f'="You have logged "&TEXT(SUM({E_}),"#,##0")&" hours so far this year across "&{WKS}&" weeks."',
    f'=IF(SUM({E_})=0,"Log days on the 📅 Daily Tracker to unlock your Year in Review.",'
    f'"🏆  Your year so far is defined by: "&INDEX({B_},MATCH(MAX({E_}),{E_},0))&"  ("&TEXT(MAX({E_}),"#,##0")&" h logged)")',
    f'="😴  At your ideal pace you will sleep "&ROUND({SD("SLEEP")}*{C_YEAR}/24,0)&" full days this year."',
    f'="💼  Work + deep work + meetings project to "&TEXT(({SD("WORK")}+{SD("DEEP")}+{SD("MEET")})*{C_YEAR},"#,##0")&" hours this year."',
    f'="✅  Free, unallocated time compounds to "&TEXT(MAX(0,{FREE})*{C_YEAR},"#,##0")&" hours a year. Spend them on purpose."',
]
for i, f in enumerate(year_lines):
    ar.merge_cells(start_row=ARL + 3 + i, start_column=2, end_row=ARL + 3 + i, end_column=9)
    put(ar, ARL + 3 + i, 2, f, F_MUTED)
ar.freeze_panes = 'A10'

# ===========================================================================
# 📊 DASHBOARD
# ===========================================================================
db = sheet(DASH, TEAL_DARK, [2, 13, 13, 13, 13, 13, 3, 12, 12, 12, 3, 26, 10])
row = kpi_strip(db, 'ASCENSION OS™ — Executive Dashboard',
    'Read-only command centre. Inputs live on 🌟 Ideal Week, 📆 Current Week and 📅 Daily Tracker.')
put(db, 3, 12, f'="Week "&WEEKNUM(TODAY(),21)&" · "&TEXT(TODAY(),"ddd d mmm yyyy")',
    Font(name='Arial', size=9, color=FAINT), RIGHT)

band(db, 8, 2, 6, '  🚦  SIGNALS')
TOTW = f'({SD("WORK")}+{SD("DEEP")}+{SD("MEET")})'
signals = [
    f'=IF({TOTW}>60,"🔴  Total work is over 60 h/week — burnout territory.",IF({TOTW}>55,"🟠  Total work is above 55 h/week — watch the load.","🟢  Work load is sustainable."))',
    f'=IF({SD("SLEEP")}<{G_SG}*7,"🔴  Sleep is below your nightly goal — recovery first.","🟢  Sleep is on target.")',
    f'=IF({SD("FIT")}<3,"🟡  Fitness is under 3 h/week — schedule workouts first.","🟢  Fitness habit is funded.")',
    f'=IF({FREE}<0,"🔴  You allocated MORE than 168 hours — something must give.",IF({FREE}<5,"🟡  Under 5 h of free time — build in slack.","🟢  Healthy buffer of free time."))',
    f'=IF({SD("DEEP")}<5,"🟡  Deep work is under 5 h/week — protect focus blocks.","🟢  Deep-work practice is funded.")',
]
for i, f in enumerate(signals):
    db.merge_cells(start_row=9 + i, start_column=2, end_row=9 + i, end_column=6)
    put(db, 9 + i, 2, f, F_BODY)

band(db, 15, 2, 6, '  🎯  GOAL PROGRESS')
header_cells(db, 16, 2, ['Goal', 'Now', 'Target', '%', ''])
goal_rows = [
    ('Productive hours', f'={PROD}', f'={G_WT}', f'={PROD}/{G_WT}'),
    ('Sleep (h/week)', f'={SD("SLEEP")}', f'={G_SG}*7', f'={SD("SLEEP")}/({G_SG}*7)'),
    ('Free time buffer', f'=MAX(0,{FREE})', 5, f'=MAX(0,{FREE})/5'),
]
for i, (lab, now, tgt, pct) in enumerate(goal_rows):
    rr = 17 + i
    put(db, rr, 2, lab, F_MUTED, border=BORDER)
    put(db, rr, 3, now, F_BOLD, CENTER, '0.0', border=BORDER)
    put(db, rr, 4, tgt, F_MUTED, CENTER, '0.0', border=BORDER)
    put(db, rr, 5, pct, F_BOLD, CENTER, '0%', border=BORDER)
    put(db, rr, 6, None, F_BODY, border=BORDER)
db.conditional_formatting.add('E17:E19',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1.5,
                color=TEAL, showValue=True))

band(db, 21, 2, 6, '  📈  TRENDS — LOGGED HOURS, LAST 4 WEEKS')
labels = ['3 weeks ago', '2 weeks ago', 'Last week', 'This week']
for i, lab in enumerate(labels):
    off = (3 - i) * 7
    put(db, 22, 2 + i, lab, Font(name='Arial', size=8, color=FAINT), CENTER)
    put(db, 23, 2 + i,
        f'=SUMIFS({LOG_ACT},{LOG_DATE},">="&({WS_CELL}-{off}),{LOG_DATE},"<"&({WS_CELL}-{off}+7))',
        F_BOLD, CENTER, '0.0')
db.conditional_formatting.add('B23:E23',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=168,
                color=TEAL, showValue=True))

band(db, 25, 2, 6, '  💡  EXECUTIVE INSIGHTS')
D_ = f'{ST}$D${M_FIRST}:$D${M_LAST}'
C_ = f'{ST}$C${M_FIRST}:$C${M_LAST}'
ins = [
    f'="🏆  Largest planned investment: "&INDEX({C_},MATCH(MAX({D_}),{D_},0))&"  ("&TEXT(MAX({D_}),"0.0")&" h/week)"',
    f'=IF(SUM({CW_ACT})=0,"📅  Log your first day on the Daily Tracker to unlock live insights.",'
    f'IF(MAX({CW_VAR})<=0,"✅  No category is over plan this week.",'
    f'"🔺  Most over plan this week: "&INDEX({C_},MATCH(MAX({CW_VAR}),{CW_VAR},0))&"  (+"&TEXT(MAX({CW_VAR}),"0.0")&" h)"))',
    f'=IF(SUM({CW_ACT})=0,"","🔻  Most behind plan this week: "&INDEX({C_},MATCH(MIN({CW_VAR}),{CW_VAR},0))&"  ("&TEXT(MIN({CW_VAR}),"0.0")&" h)")',
]
for i, f in enumerate(ins):
    db.merge_cells(start_row=26 + i, start_column=2, end_row=26 + i, end_column=6)
    put(db, 26 + i, 2, f, F_MUTED)

# life-area table (chart source) — core rows + Free Time contiguous
LT = 8
put(db, LT, 12, 'LIFE AREA', F_HEAD, LEFT, fill=FILL_TEAL, border=BORDER)
put(db, LT, 13, 'HOURS', F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
for i, (key, emoji, name, *_rest) in enumerate(CATS):
    rr = LT + 1 + i
    put(db, rr, 12, f'{emoji}  {name}', F_BODY, border=BORDER)
    put(db, rr, 13, f'={ST}D{M_FIRST + i}', F_BODY, CENTER, '0.0', border=BORDER)
FREE_ROW = LT + 1 + N_CORE
fc = put(db, FREE_ROW, 12, '✅  Free Time', F_BOLD, border=BORDER)
fv = put(db, FREE_ROW, 13, f'=MAX(0,{FREE})', F_BOLD, CENTER, '0.0', border=BORDER)
fc.font = fv.font = Font(name='Arial', size=10, bold=True, color=GREEN)
for i in range(N_CUSTOM):
    rr = FREE_ROW + 1 + i
    sr = M_FIRST + N_CORE + i
    put(db, rr, 12, f'=IF({ST}C{sr}="","",{ST}B{sr}&"  "&{ST}C{sr})', F_MUTED, border=BORDER)
    put(db, rr, 13, f'=IF({ST}C{sr}="","",{ST}D{sr})', F_MUTED, CENTER, '0.0', border=BORDER)
tr_ = FREE_ROW + N_CUSTOM + 1
tc = put(db, tr_, 12, '❌  Total Non-Free Time', F_BOLD, border=BORDER)
tv = put(db, tr_, 13, f'={USED}', F_BOLD, CENTER, '0.0', border=BORDER)
tc.font = tv.font = Font(name='Arial', size=10, bold=True, color=RED)

data = Reference(db, min_col=13, min_row=LT, max_row=FREE_ROW)
cats = Reference(db, min_col=12, min_row=LT + 1, max_row=FREE_ROW)
pie = PieChart()
pie.title = 'What does my time split look like by %?'
pie.add_data(data, titles_from_data=True); pie.set_categories(cats)
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
pie.height, pie.width = 8.2, 10.2
db.add_chart(pie, 'B31')
bar = BarChart(); bar.type = 'bar'
bar.title = 'What does my time split look like overall?'
bar.add_data(data, titles_from_data=True); bar.set_categories(cats)
bar.legend = None
bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True; bar.dataLabels.numFmt = '0.0'
bar.height, bar.width = 10.5, 10.2
s = bar.series[0]; s.graphicalProperties.solidFill = ORANGE
db.add_chart(bar, 'G31')

# ===========================================================================
# 📖 START HERE
# ===========================================================================
gd = sheet(GUI, 'B7B7B7', [2, 4, 110])
put(gd, 2, 2, '📖 Start Here — how to run ASCENSION OS™', F_TITLE)
sections = [
    ('🧭  THE OPERATING SYSTEM', [
        'Everyone gets 24 hours a day and 168 hours a week. ASCENSION OS™ treats them like a budget: design on 🌟 Ideal Week, commit on 📆 Current Week, execute on 📅 Daily Tracker — every dashboard and review reconciles automatically.',
        'Every sheet carries the same executive KPI strip on top, plus a live line of what you logged today / this week / this month / this year — no tab switching needed to know where you stand.']),
    ('✏️  RULE #1 — ONLY EDIT LIGHT-BLUE CELLS', [
        'Every editable cell has a light-blue background. Everything else is a formula.']),
    ('🚀  THE OPERATING LOOP', [
        '1.  DESIGN — answer the guided sentences on 🌟 Ideal Week.',
        '2.  COMMIT — on 📆 Current Week, optionally override any category for this week ("My plan"). Blank = ideal.',
        '3.  EXECUTE — log real blocks on 📅 Daily Tracker: date, category (dropdown), hours, notes. The Today panel scores your day live.',
        '4.  REVIEW — 🗓️ Monthly Review and 🎯 Annual Review build themselves. 📊 Dashboard is the command centre.']),
    ('➕  CUSTOM CATEGORIES & GOALS', [
        'Add categories in the blue rows at the bottom of the ⚙️ Settings master table.',
        'Work days, sleep goal and your weekly productive target also live on ⚙️ Settings.']),
    ('🎯  BALANCE SCORE', [
        'A 0–100 gauge: sleep at goal, total work (work + deep work + meetings) ≤ 55 h, fitness ≥ 3 h, relationships ≥ 3 h, learning + personal development ≥ 2 h, free time ≥ 5 h, nothing over-allocated. 80+ = sustainable.']),
    ('📐  ASSUMPTIONS', [
        'Monthly = weekly × 4, yearly = weekly × 52 (editable on ⚙️ Settings; use 4.345 / 52.18 for calendar-accurate maths).',
        'The week runs Monday–Sunday. The Daily Tracker holds 150 log rows — clear old rows when you start a fresh period.']),
]
gr = 4
for title, paras in sections:
    band(gd, gr, 2, 3, '  ' + title); gr += 1
    for p in paras:
        c = put(gd, gr, 3, p, F_MUTED)
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        gd.row_dimensions[gr].height = max(16, 15 * (1 + len(p) // 105))
        gr += 1
    gr += 1

wb.move_sheet(DASH, offset=-(len(wb.sheetnames) - 1))
wb.active = wb[DASH]
wb.save(OUT)
print('saved', OUT)

import subprocess
res = subprocess.run(['python3', '/root/.claude/skills/xlsx/scripts/recalc.py', OUT, '240'],
                     capture_output=True, text=True)
print(res.stdout[-1500:] or res.stderr[-1500:])
