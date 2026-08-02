#!/usr/bin/env python3
"""Build the 168-Hour Executive OS as a ready-to-import .xlsx.

Output: 168-Hour-Executive-OS.xlsx (repo root)
Import into Google Sheets: Drive -> New -> File upload -> open with Google Sheets
(or in an existing sheet: File -> Import -> Upload).

Design follows the user's "Ascension OS Weekly Tracker" reference:
teal section bands, light-blue input cells, guided sentences with
"This means I spend ... hours per week ..." columns, top summary
(WEEKLY / MONTHLY / ANNUALLY x HOURS USED / HOURS FREE), LIFE AREA/HOURS
panel, pie + orange bar charts. Default values match the reference
screenshots (used 166.1 h, free 1.9 h).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = '/home/user/Team/168-Hour-Executive-OS.xlsx'

# ---- palette (sampled from the reference screenshots) -----------------------
TEAL       = '2BB5AE'   # section bands / table headers
TEAL_DARK  = '17817B'
BLUE_BOX   = 'A9C7E7'   # "Current Week" hero box
INPUT_BG   = 'DDEBF7'   # editable cells
INPUT_LN   = '9DC3E6'
INK        = '1F2937'
SLATE      = '5B6B7C'
ORANGE     = 'F6A21D'   # bar chart, like reference
GREEN      = '2E9E4F'
GREEN_BG   = 'D9F2E3'
RED        = 'C0392B'
RED_BG     = 'FADBD8'
AMBER      = 'B7791F'
AMBER_BG   = 'FCEFC7'
LINE       = 'BFC9D4'
CARD       = 'F3F6F9'

DASH, INP, DAY, WKR = '📊 Dashboard', '📝 Weekly Tracker', '📅 Daily Tracker', '📈 Weekly Review'
MON, ANN, SET, GUI = '🗓️ Monthly Tracker', '🎯 Annual Tracker', '⚙️ Settings', '📖 Start Here'

F_BODY  = Font(name='Arial', size=10, color=INK)
F_MUTED = Font(name='Arial', size=10, color=SLATE)
F_BOLD  = Font(name='Arial', size=10, color=INK, bold=True)
F_BAND  = Font(name='Arial', size=11, color='FFFFFF', bold=True)
F_HEAD  = Font(name='Arial', size=9,  color='FFFFFF', bold=True)
F_TITLE = Font(name='Arial', size=16, color=INK, bold=True)
F_SUB   = Font(name='Arial', size=10, color=SLATE, italic=True)
F_KPI   = Font(name='Arial', size=20, color=INK, bold=True)

FILL_TEAL  = PatternFill('solid', fgColor=TEAL)
FILL_INPUT = PatternFill('solid', fgColor=INPUT_BG)
FILL_BOX   = PatternFill('solid', fgColor=BLUE_BOX)
FILL_CARD  = PatternFill('solid', fgColor=CARD)

thin = Side(style='thin', color=LINE)
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
in_side = Side(style='thin', color=INPUT_LN)
BORDER_IN = Border(top=in_side, bottom=in_side, left=in_side, right=in_side)

CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left',  vertical='center')
RIGHT  = Alignment(horizontal='right', vertical='center')
WRAP   = Alignment(horizontal='left',  vertical='center', wrap_text=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sheet(name, tab, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def band(ws, row, c1, c2, text):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row=row, column=c1, value=text)
    c.font, c.alignment = F_BAND, LEFT
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = FILL_TEAL
    ws.row_dimensions[row].height = 22


def put(ws, row, col, value, font=F_BODY, align=LEFT, fmt=None, fill=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment = font, align
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = border
    return c


def input_cell(ws, row, col, value, fmt='0.##'):
    return put(ws, row, col, value, font=F_BOLD, align=CENTER, fmt=fmt,
               fill=FILL_INPUT, border=BORDER_IN)


# ============================================================================
# 📝 WEEKLY TRACKER (inputs) — reference-style guided sentences
# cols: A pad | B sentence | C input | D unit | E "This means I spend" | F val | G tail
# ============================================================================
inp = sheet(INP, TEAL, [2, 30, 9, 40, 21, 9, 44])
put(inp, 2, 2, '📝 Weekly Tracker — plan your 168 hours', F_TITLE)
put(inp, 3, 2, 'Only edit the LIGHT-BLUE cells. Every other number updates automatically — here, on the Dashboard and in every tracker.', F_SUB)

r = 5
plan = {}     # key -> (sheet, 'F12') total cell ref
work_days = None

def line(label, val, unit, calc=None, tail=None, note=None, vfmt='0.##'):
    """One sentence row; calc is a formula factory given current row."""
    global r
    put(inp, r, 2, label)
    input_cell(inp, r, 3, val, vfmt)
    put(inp, r, 4, unit, F_MUTED)
    if calc:
        put(inp, r, 5, 'This means I spend', F_MUTED, RIGHT)
        put(inp, r, 6, calc(r), F_BOLD, CENTER, '0.0')
        put(inp, r, 7, tail, F_MUTED)
    elif note:
        put(inp, r, 7, note, Font(name='Arial', size=8, color=SLATE, italic=True))
    r += 1
    return r - 1

def calc_rows(formula, tail):
    global r
    put(inp, r, 5, 'Which means I spend', F_MUTED, RIGHT)
    put(inp, r, 6, formula, F_BOLD, CENTER, '0.0')
    put(inp, r, 7, tail, F_MUTED)
    r += 1
    return r - 1

sections = []   # (key, emoji, name, total_row)
prev_leaves = None

def open_section(emoji, name):
    global r
    band(inp, r, 2, 7, f'  {emoji}  {name}')
    r += 1

def close_section(key, emoji, name, f_rows, total_tail):
    """Total + cascading 'that leaves me with' rows."""
    global r, prev_leaves
    put(inp, r, 5, f'{emoji} In total, I spend', F_BOLD, RIGHT)
    tf = put(inp, r, 6, '=' + '+'.join(f'F{x}' for x in f_rows), F_BOLD, CENTER, '0.0')
    put(inp, r, 7, total_tail, F_BOLD)
    for col in range(2, 8):
        inp.cell(row=r, column=col).border = Border(top=Side(style='thin', color=TEAL))
    plan[key] = f'F{r}'
    total_row = r
    r += 1
    put(inp, r, 5, '✅ That leaves me with', F_BOLD, RIGHT)
    base = '168' if prev_leaves is None else f'F{prev_leaves}'
    lv = put(inp, r, 6, f'={base}-F{total_row}', F_BOLD, CENTER, '0.0')
    lv.font = Font(name='Arial', size=10, bold=True, color=GREEN)
    put(inp, r, 7, 'hours per week for everything else.', F_MUTED)
    prev_leaves = r
    r += 2
    sections.append((key, emoji, name, total_row))

# ---- SLEEP ------------------------------------------------------------------
open_section('😴', 'SLEEP')
a = line('I sleep for around', 8, 'hours per night.',
         lambda rr: f'=C{rr}*7', 'hours per week sleeping.')
b = line('I wind down for', 10, 'minutes before sleeping.',
         lambda rr: f'=C{rr}*7/60', 'hours per week preparing to sleep.')
c = line('My wake-up routine takes', 0, 'minutes each morning.',
         lambda rr: f'=C{rr}*7/60', 'hours per week waking up.')
close_section('SLEEP', '😴', 'Sleep', [a, b, c], 'hours per week on sleep-related things.')

# ---- WORK -------------------------------------------------------------------
open_section('💻', 'WORK')
h_day = line('I work for around', 8, 'hours per day (excluding lunch break).')
d_row = line('I work', 6, 'days per week.',
             lambda rr: f'=C{h_day}*C{rr}', 'hours per week officially working.', vfmt='0')
work_days = f"'{INP}'!$C${d_row}"
lunch = line('My lunch break is', 30, 'minutes per working day.',
             lambda rr: f'=C{rr}*C{d_row}/60', 'hours per week on my lunch break.')
comm = line('My commute is', 0, 'minutes each way.',
            lambda rr: f'=C{rr}*2*C{d_row}/60', 'hours per week commuting.')
prep = line('I spend', 10, 'minutes per day getting ready for work.',
            lambda rr: f'=C{rr}*C{d_row}/60', 'hours per week getting ready for work.')
chg = line('I spend', 10, 'minutes per day getting changed after work.',
           lambda rr: f'=C{rr}*C{d_row}/60', 'hours per week getting changed after work.')
line('… of which meetings take', 5, 'hours per week.',
     note='already inside your working hours — not added again')
line('… and deep work takes', 12, 'hours per week.',
     note='already inside your working hours — not added again')
close_section('WORK', '💻', 'Work', [d_row, lunch, comm, prep, chg],
              'hours per week working.')

# ---- FOOD -------------------------------------------------------------------
open_section('🍕', 'FOOD')
b1 = line('Each day, I spend around', 15, 'minutes eating breakfast.')
b2 = line('And another', 0, 'minutes preparing/cooking/ordering it.')
fb = calc_rows(f'=(C{b1}+C{b2})*7/60', 'hours per week on breakfast.')
l1 = line('Each non-work day, I spend', 40, 'minutes eating lunch.')
l2 = line('And another', 0, 'minutes preparing/cooking/ordering it.')
fl = calc_rows(f'=(C{l1}+C{l2})*(7-C{d_row})/60', 'hours per week on lunch (excluding work days).')
n1 = line('Each day, I spend around', 30, 'minutes eating dinner.')
n2 = line('And another', 0, 'minutes preparing/cooking/ordering it.')
fd = calc_rows(f'=(C{n1}+C{n2})*7/60', 'hours per week on dinner.')
close_section('FOOD', '🍕', 'Food', [fb, fl, fd], 'hours per week on cooking +/- eating food.')

# ---- HOUSEHOLD CHORES -------------------------------------------------------
open_section('🧹', 'HOUSEHOLD CHORES')
g1 = line('I do the groceries', 2, 'times per week.', vfmt='0')
g2 = line('Each grocery session takes', 120, 'minutes (including travel, loading, unloading etc).',
          lambda rr: f'=C{g1}*C{rr}/60', 'hours per week on groceries.')
c1r = line('I clean the house approx.', 1, 'times per week.', vfmt='0')
c2r = line('Each cleaning session takes', 60, 'minutes.',
           lambda rr: f'=C{c1r}*C{rr}/60', 'hours per week on cleaning.')
w1 = line('I do laundry', 0, 'times per week.', vfmt='0')
w2 = line('Each laundry session takes me', 0, 'minutes.',
          lambda rr: f'=C{w1}*C{rr}/60', 'hours per week on laundry.')
m1 = line('I do other misc. chores', 1, 'hour(s) per week.',
          lambda rr: f'=C{rr}', 'hours per week on misc. chores.')
close_section('CHORES', '🧹', 'Household Chores', [g2, c2r, w2, m1], 'hours per week on chores.')

# ---- CHILDCARE --------------------------------------------------------------
open_section('👶', 'CHILDCARE')
k1 = line('Each weekday, I spend', 2, 'hours actively taking care of my child(ren).',
          lambda rr: f'=C{rr}*5', 'hours on weekdays on childcare.')
k2 = line('Each weekend day, I spend', 3, 'hours actively taking care of my child(ren).',
          lambda rr: f'=C{rr}*2', 'hours on weekends on childcare.')
k3 = line('Each week, I spend', 0, 'hours on child admin & logistics etc.',
          lambda rr: f'=C{rr}', 'hours per week on child admin.')
k4 = line('Each week, I spend', 0, 'hours on other child-themed stuff.',
          lambda rr: f'=C{rr}', 'hours per week on other child time.')
close_section('CHILD', '👶', 'Childcare', [k1, k2, k3, k4], 'hours per week on childcare.')

# ---- FITNESS ----------------------------------------------------------------
open_section('🏋️', 'FITNESS')
f1 = line('Each week, I spend', 8, 'hours exercising (including commute time).',
          lambda rr: f'=C{rr}', 'hours per week exercising.')
f2 = line('Each week, I spend', 2, 'hours on other fitness-related things.',
          lambda rr: f'=C{rr}', 'hours per week on other fitness things.')
close_section('FIT', '🏋️', 'Fitness', [f1, f2], 'hours per week on fitness-related stuff.')

# ---- ENTERTAINMENT ----------------------------------------------------------
open_section('📺', 'ENTERTAINMENT')
e1 = line('Each week, I spend', 0, 'hours watching TV shows / movies.', lambda rr: f'=C{rr}', 'hours per week on TV / movies.')
e2 = line('Each week, I spend', 7, 'hours on social media apps.', lambda rr: f'=C{rr}', 'hours per week on social media.')
e3 = line('Each week, I spend', 2, 'hours reading.', lambda rr: f'=C{rr}', 'hours per week reading.')
e4 = line('Each week, I spend', 0, 'hours gaming.', lambda rr: f'=C{rr}', 'hours per week gaming.')
e5 = line('Each week, I spend', 3, 'hours on other "entertainment" stuff.', lambda rr: f'=C{rr}', 'hours per week on other entertainment.')
close_section('ENT', '📺', 'Entertainment', [e1, e2, e3, e4, e5], 'hours per week on entertainment.')

# ---- RELATIONSHIPS ----------------------------------------------------------
open_section('👨‍👩‍👧', 'RELATIONSHIPS')
r1 = line('Each week, I spend', 3, 'hours on quality family time.', lambda rr: f'=C{rr}', 'hours per week on family time.')
r2 = line('Each week, I spend', 3, 'hours on other socialising time.', lambda rr: f'=C{rr}', 'hours per week socialising.')
close_section('REL', '👨‍👩‍👧', 'Relationships', [r1, r2], 'hours per week on relationships.')

# ---- LEARNING ---------------------------------------------------------------
open_section('📚', 'LEARNING')
q1 = line('Each week, I spend', 0, 'hours on courses & structured learning.', lambda rr: f'=C{rr}', 'hours per week on courses.')
q2 = line('Each week, I spend', 0, 'hours on books, articles & study.', lambda rr: f'=C{rr}', 'hours per week studying.')
q3 = line('Each day, I listen to', 0, 'minutes of podcasts / audiobooks.', lambda rr: f'=C{rr}*7/60', 'hours per week on podcasts.')
close_section('LEARN', '📚', 'Learning', [q1, q2, q3], 'hours per week on learning.')

# ---- PERSONAL DEVELOPMENT ---------------------------------------------------
open_section('🌱', 'PERSONAL DEVELOPMENT')
p1 = line('Each day, I journal for', 0, 'minutes.', lambda rr: f'=C{rr}*7/60', 'hours per week journaling.')
p2 = line('Each day, I meditate for', 0, 'minutes.', lambda rr: f'=C{rr}*7/60', 'hours per week meditating.')
p3 = line('Each week, I spend', 0, 'hours on planning & weekly review.', lambda rr: f'=C{rr}', 'hours per week planning.')
close_section('PD', '🌱', 'Personal Development', [p1, p2, p3], 'hours per week on personal development.')

# ---- SPIRITUAL --------------------------------------------------------------
open_section('🙏', 'SPIRITUAL')
s1 = line('Each day, I spend', 0, 'minutes on spiritual practice.', lambda rr: f'=C{rr}*7/60', 'hours per week on daily practice.')
s2 = line('Each week, I spend', 0, 'hours on services & community.', lambda rr: f'=C{rr}', 'hours per week on services.')
close_section('SPIRIT', '🙏', 'Spiritual', [s1, s2], 'hours per week on spiritual life.')

# ---- TRAVEL -----------------------------------------------------------------
open_section('🚗', 'TRAVEL')
t1 = line('Each week, I spend', 0, 'hours on errands & getting around (non-commute).', lambda rr: f'=C{rr}', 'hours per week on errands & travel.')
t2 = line('Each week, I spend', 0, 'hours on trips & other travel.', lambda rr: f'=C{rr}', 'hours per week on trips.')
close_section('TRAVEL', '🚗', 'Travel', [t1, t2], 'hours per week travelling.')

# ---- ADMIN ------------------------------------------------------------------
open_section('🗂️', 'ADMIN')
a1 = line('Each week, I spend', 0, 'hours on finances, email & life admin.', lambda rr: f'=C{rr}', 'hours per week on life admin.')
a2 = line('Each week, I spend', 0, 'minutes on appointments & paperwork.', lambda rr: f'=C{rr}/60', 'hours per week on paperwork.')
close_section('ADMIN', '🗂️', 'Admin', [a1, a2], 'hours per week on admin.')

inp.freeze_panes = 'A5'
dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='10000',
                    allow_blank=True, showErrorMessage=True,
                    error='Enter a number (0 or more).')
inp.add_data_validation(dv)
dv.add(f'C5:C{r}')

# ============================================================================
# ⚙️ SETTINGS — master table, constants, custom categories
# ============================================================================
st = sheet(SET, TEAL_DARK, [2, 5, 24, 11, 10, 9, 11, 9, 9, 10])
put(st, 2, 2, '⚙️ Settings — master data', F_TITLE)
put(st, 3, 2, 'One source of truth. Add your own categories in the light-blue rows — every tracker and the Dashboard pick them up automatically.', F_SUB)

band(st, 5, 2, 10, '  ⏱️  CONSTANTS')
consts = [('Hours per day', 24, 'HRS_DAY'),
          ('Hours per week', '=D6*7', 'HRS_WEEK'),
          ('Weeks per month', 4, 'WKS_MON'),
          ('Weeks per year', 52, 'WKS_YEAR')]
for i, (label, val, _) in enumerate(consts):
    rr = 6 + i
    put(st, rr, 2, label, F_MUTED)
    st.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
    put(st, rr, 4, val, F_BOLD, CENTER, '0.###')
C_WEEK  = f"'{SET}'!$D$7"
C_MONTH = f"'{SET}'!$D$8"
C_YEAR  = f"'{SET}'!$D$9"
st['D8'].comment = Comment('Executive convention: 4 weeks/month, 52 weeks/year (matches the reference sheet). For calendar-accurate maths change to 4.345 and 52.18.', 'Executive OS')

TBL = 12
band(st, TBL, 2, 10, '  🗂️  MASTER CATEGORY TABLE')
headers = ['', 'Category', 'Hours / wk', '% of week', 'Per day', 'Productive', 'Min h/wk', 'Max h/wk', 'Status']
for i, h in enumerate(headers):
    cell = put(st, TBL + 1, 2 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
FIRST = TBL + 2                       # 14
META = {  # key -> (emoji, name, productive, min, max)
    'SLEEP': ('😴', 'Sleep', False, 49, 63), 'WORK': ('💻', 'Work', True, '', 55),
    'FOOD': ('🍕', 'Food', False, 3.5, 21), 'CHORES': ('🧹', 'Household Chores', False, '', 15),
    'CHILD': ('👶', 'Childcare', False, '', ''), 'FIT': ('🏋️', 'Fitness', True, 3, 20),
    'ENT': ('📺', 'Entertainment', False, '', 25), 'REL': ('👨‍👩‍👧', 'Relationships', False, 3, ''),
    'LEARN': ('📚', 'Learning', True, 1, ''), 'PD': ('🌱', 'Personal Development', True, 1, ''),
    'SPIRIT': ('🙏', 'Spiritual', False, '', ''), 'TRAVEL': ('🚗', 'Travel', False, '', ''),
    'ADMIN': ('🗂️', 'Admin', True, '', ''),
}
ORDER = ['SLEEP', 'WORK', 'FOOD', 'CHORES', 'CHILD', 'FIT', 'ENT', 'REL',
         'LEARN', 'PD', 'SPIRIT', 'TRAVEL', 'ADMIN']
N_CORE, N_CUSTOM = len(ORDER), 5
LAST = FIRST + N_CORE + N_CUSTOM - 1   # 31

for i, key in enumerate(ORDER):
    rr = FIRST + i
    emoji, name, prod, mn, mx = META[key]
    put(st, rr, 2, emoji, F_BODY, CENTER, border=BORDER)
    put(st, rr, 3, name, F_BODY, border=BORDER)
    put(st, rr, 4, f"='{INP}'!{plan[key]}", F_BOLD, CENTER, '0.0', border=BORDER)
    put(st, rr, 7, prod, F_BODY, CENTER, border=BORDER)
    put(st, rr, 8, mn if mn != '' else None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 9, mx if mx != '' else None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
for i in range(N_CUSTOM):
    rr = FIRST + N_CORE + i
    put(st, rr, 2, None, F_BODY, CENTER, fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 3, None, F_BODY, fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 4, None, F_BOLD, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 7, False, F_BODY, CENTER, fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 8, None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
    put(st, rr, 9, None, F_BODY, CENTER, '0.0', fill=FILL_INPUT, border=BORDER_IN)
st.cell(row=FIRST + N_CORE, column=3).comment = Comment(
    'Type a custom category name (e.g. Content Creation) + weekly hours. '
    'It flows into every tracker automatically.', 'Executive OS')
for rr in range(FIRST, LAST + 1):
    put(st, rr, 5, f'=IF($C{rr}="","",$D{rr}/{C_WEEK})', F_BODY, CENTER, '0.0%', border=BORDER)
    put(st, rr, 6, f'=IF($C{rr}="","",$D{rr}/7)', F_BODY, CENTER, '0.0', border=BORDER)
    put(st, rr, 10,
        f'=IF($C{rr}="","",IF(AND($H{rr}<>"",$D{rr}<$H{rr}),"🔻 Low",'
        f'IF(AND($I{rr}<>"",$D{rr}>$I{rr}),"🔺 High","✅ OK")))',
        F_BODY, CENTER, border=BORDER)

TOT = LAST + 2
put(st, TOT, 3, 'Total allocated', F_BOLD, RIGHT)
put(st, TOT, 4, f'=SUM(D{FIRST}:D{LAST})', F_BOLD, CENTER, '0.0')
put(st, TOT + 1, 3, 'Free time (168 − allocated)', F_BOLD, RIGHT)
put(st, TOT + 1, 4, f'={C_WEEK}-D{TOT}', F_BOLD, CENTER, '0.0')
put(st, TOT + 2, 3, 'Productive hours', F_BOLD, RIGHT)
put(st, TOT + 2, 4, f'=SUMPRODUCT(--(G{FIRST}:G{LAST}=TRUE),D{FIRST}:D{LAST})',
    F_BOLD, CENTER, '0.0')
USED = f"'{SET}'!$D${TOT}"
FREE = f"'{SET}'!$D${TOT + 1}"
PROD = f"'{SET}'!$D${TOT + 2}"

st.conditional_formatting.add(
    f'D{FIRST}:D{LAST}',
    FormulaRule(formula=[f'AND($I{FIRST}<>"",$D{FIRST}>$I{FIRST})'],
                fill=PatternFill('solid', fgColor=RED_BG),
                font=Font(name='Arial', size=10, color=RED, bold=True)))
st.conditional_formatting.add(
    f'D{FIRST}:D{LAST}',
    FormulaRule(formula=[f'AND($H{FIRST}<>"",$D{FIRST}<$H{FIRST})'],
                fill=PatternFill('solid', fgColor=AMBER_BG),
                font=Font(name='Arial', size=10, color=AMBER, bold=True)))
st.conditional_formatting.add(
    f'D{TOT + 1}',
    FormulaRule(formula=[f'$D${TOT + 1}<0'], fill=PatternFill('solid', fgColor=RED_BG),
                font=Font(name='Arial', size=10, color=RED, bold=True)))
st.freeze_panes = f'A{TBL + 1}'

def cat_label(rr):
    return f"=IF('{SET}'!C{rr}=\"\",\"\",'{SET}'!B{rr}&\"  \"&'{SET}'!C{rr})"

# ============================================================================
# 📅 DAILY TRACKER
# ============================================================================
dy = sheet(DAY, TEAL, [2, 26, 9, 7, 7, 7, 7, 7, 7, 7, 10, 9])
put(dy, 2, 2, '📅 Daily Tracker — log what actually happened', F_TITLE)
put(dy, 3, 2, 'Type real hours in the light-blue grid. Categories flow in from ⚙️ Settings automatically.', F_SUB)
HEAD = 5
cols = ['Category', 'Plan/day', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun', 'Week', 'vs plan']
for i, h in enumerate(cols):
    put(dy, HEAD, 2 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
D1 = HEAD + 1
for i in range(N_CORE + N_CUSTOM):
    rr, sr = D1 + i, FIRST + i
    put(dy, rr, 2, cat_label(sr), F_BODY, border=BORDER)
    put(dy, rr, 3, f"=IF($B{rr}=\"\",\"\",'{SET}'!D{sr}/7)", F_MUTED, CENTER, '0.00', border=BORDER)
    for cno in range(4, 11):
        put(dy, rr, cno, None, F_BODY, CENTER, '0.0#', fill=FILL_INPUT, border=BORDER_IN)
    put(dy, rr, 11, f'=IF($B{rr}="","",SUM(D{rr}:J{rr}))', F_BOLD, CENTER, '0.0', border=BORDER)
    put(dy, rr, 12, f"=IF($B{rr}=\"\",\"\",K{rr}-'{SET}'!D{sr})", F_BODY, CENTER, '0.0', border=BORDER)
DL = D1 + N_CORE + N_CUSTOM - 1
TR = DL + 1
put(dy, TR, 2, 'Hours logged', F_BOLD, RIGHT)
put(dy, TR + 1, 2, 'Unlogged (24 − logged)', F_MUTED, RIGHT)
for cno in range(4, 11):
    L = get_column_letter(cno)
    put(dy, TR, cno, f'=SUM({L}{D1}:{L}{DL})', F_BOLD, CENTER, '0.0')
    put(dy, TR + 1, cno, f'=24-{L}{TR}', F_MUTED, CENTER, '0.0')
put(dy, TR, 11, f'=SUM(K{D1}:K{DL})', F_BOLD, CENTER, '0.0')
put(dy, TR + 1, 11, f'={C_WEEK}-K{TR}', F_MUTED, CENTER, '0.0')
dy.conditional_formatting.add(
    f'D{TR}:J{TR}',
    FormulaRule(formula=[f'D{TR}>24'], fill=PatternFill('solid', fgColor=RED_BG),
                font=Font(name='Arial', size=10, color=RED, bold=True)))
dy.conditional_formatting.add(
    f'D{TR}:J{TR}',
    FormulaRule(formula=[f'D{TR}>18'], fill=PatternFill('solid', fgColor=AMBER_BG),
                font=Font(name='Arial', size=10, color=AMBER, bold=True)))
dv2 = DataValidation(type='decimal', operator='between', formula1='0', formula2='24',
                     allow_blank=True, showErrorMessage=True, error='Hours must be 0–24.')
dy.add_data_validation(dv2)
dv2.add(f'D{D1}:J{DL}')
dy.freeze_panes = f'A{D1}'

# ============================================================================
# 📈 WEEKLY REVIEW — planned vs actual
# ============================================================================
wk = sheet(WKR, TEAL, [2, 26, 10, 10, 9, 10, 10])
put(wk, 2, 2, '📈 Weekly Review — planned vs actual', F_TITLE)
put(wk, 3, 2, 'Planned comes from 📝 Weekly Tracker; actual from 📅 Daily Tracker. Nothing to type here.', F_SUB)
WH = 5
for i, h in enumerate(['Category', 'Planned', 'Actual', 'Diff', '% of plan', '% of week']):
    put(wk, WH, 2 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
W1 = WH + 1
for i in range(N_CORE + N_CUSTOM):
    rr, sr, drr = W1 + i, FIRST + i, D1 + i
    put(wk, rr, 2, cat_label(sr), F_BODY, border=BORDER)
    put(wk, rr, 3, f"=IF($B{rr}=\"\",\"\",'{SET}'!D{sr})", F_BODY, CENTER, '0.0', border=BORDER)
    put(wk, rr, 4, f"=IF($B{rr}=\"\",\"\",'{DAY}'!K{drr})", F_BOLD, CENTER, '0.0', border=BORDER)
    put(wk, rr, 5, f'=IF($B{rr}="","",D{rr}-C{rr})', F_BODY, CENTER, '0.0', border=BORDER)
    put(wk, rr, 6, f'=IF(OR($B{rr}="",C{rr}=0),"",D{rr}/C{rr})', F_BODY, CENTER, '0%', border=BORDER)
    put(wk, rr, 7, f'=IF($B{rr}="","",C{rr}/{C_WEEK})', F_MUTED, CENTER, '0.0%', border=BORDER)
WL = W1 + N_CORE + N_CUSTOM - 1
WT = WL + 2
put(wk, WT, 2, 'Total allocated / logged', F_BOLD, RIGHT)
put(wk, WT, 3, f'=SUM(C{W1}:C{WL})', F_BOLD, CENTER, '0.0')
put(wk, WT, 4, f'=SUM(D{W1}:D{WL})', F_BOLD, CENTER, '0.0')
put(wk, WT + 1, 2, 'Free / remaining hours', F_BOLD, RIGHT)
put(wk, WT + 1, 3, f'={C_WEEK}-C{WT}', F_BOLD, CENTER, '0.0')
put(wk, WT + 1, 4, f'={C_WEEK}-D{WT}', F_BOLD, CENTER, '0.0')
put(wk, WT + 2, 2, 'Utilization', F_BOLD, RIGHT)
put(wk, WT + 2, 3, f'=C{WT}/{C_WEEK}', F_BOLD, CENTER, '0.0%')
put(wk, WT + 2, 4, f'=D{WT}/{C_WEEK}', F_BOLD, CENTER, '0.0%')
wk.conditional_formatting.add(
    f'D{W1}:D{WL}',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=60,
                color=ORANGE, showValue=True))
wk.conditional_formatting.add(
    f'F{W1}:F{WL}',
    FormulaRule(formula=[f'AND($B{W1}<>"",$F{W1}<>"",$F{W1}>1.2)'],
                fill=PatternFill('solid', fgColor=RED_BG),
                font=Font(name='Arial', size=10, color=RED)))
wk.conditional_formatting.add(
    f'F{W1}:F{WL}',
    FormulaRule(formula=[f'AND($B{W1}<>"",$F{W1}<>"",$F{W1}<0.5)'],
                fill=PatternFill('solid', fgColor=AMBER_BG),
                font=Font(name='Arial', size=10, color=AMBER)))
wk.freeze_panes = f'A{W1}'

# ============================================================================
# 🗓️ MONTHLY TRACKER
# ============================================================================
mo = sheet(MON, TEAL, [2, 26, 10, 11, 12, 10, 11])
put(mo, 2, 2, '🗓️ Monthly Tracker — the month, projected', F_TITLE)
put(mo, 3, 2, 'Weekly plan × 4 weeks (change the constant on ⚙️ Settings). Actuals extrapolate the current week.', F_SUB)
MH = 5
for i, h in enumerate(['Category', 'Weekly plan', 'Monthly plan', 'Monthly actual*', 'Variance', '% of month']):
    put(mo, MH, 2 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
M1 = MH + 1
for i in range(N_CORE + N_CUSTOM):
    rr, sr, wrr = M1 + i, FIRST + i, W1 + i
    put(mo, rr, 2, cat_label(sr), F_BODY, border=BORDER)
    put(mo, rr, 3, f"=IF($B{rr}=\"\",\"\",'{SET}'!D{sr})", F_BODY, CENTER, '0.0', border=BORDER)
    put(mo, rr, 4, f'=IF($B{rr}="","",C{rr}*{C_MONTH})', F_BOLD, CENTER, '0.0', border=BORDER)
    put(mo, rr, 5, f"=IF($B{rr}=\"\",\"\",'{WKR}'!D{wrr}*{C_MONTH})", F_BODY, CENTER, '0.0', border=BORDER)
    put(mo, rr, 6, f'=IF($B{rr}="","",E{rr}-D{rr})', F_BODY, CENTER, '0.0', border=BORDER)
    put(mo, rr, 7, f'=IF($B{rr}="","",D{rr}/({C_WEEK}*{C_MONTH}))', F_MUTED, CENTER, '0.0%', border=BORDER)
ML = M1 + N_CORE + N_CUSTOM - 1
put(mo, ML + 2, 2, 'Total', F_BOLD, RIGHT)
put(mo, ML + 2, 4, f'=SUM(D{M1}:D{ML})', F_BOLD, CENTER, '0.0')
put(mo, ML + 2, 5, f'=SUM(E{M1}:E{ML})', F_BOLD, CENTER, '0.0')
put(mo, ML + 3, 2, 'Free time per month', F_BOLD, RIGHT)
put(mo, ML + 3, 4, f'={C_WEEK}*{C_MONTH}-D{ML + 2}', F_BOLD, CENTER, '0.0')
put(mo, ML + 4, 2, '* extrapolates the current Daily Tracker week across the month.',
    Font(name='Arial', size=8, color=SLATE, italic=True))
mo.freeze_panes = f'A{M1}'

# ============================================================================
# 🎯 ANNUAL TRACKER
# ============================================================================
an = sheet(ANN, TEAL, [2, 26, 10, 12, 11, 11, 10])
put(an, 2, 2, '🎯 Annual Tracker — a year of this plan', F_TITLE)
put(an, 3, 2, 'Weekly plan × 52 weeks. This is where small weekly numbers become life-sized.', F_SUB)
AH = 5
for i, h in enumerate(['Category', 'Weekly h', 'Annual hours', 'Annual days', 'Annual weeks', '% of year']):
    put(an, AH, 2 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
A1 = AH + 1
for i in range(N_CORE + N_CUSTOM):
    rr, sr = A1 + i, FIRST + i
    put(an, rr, 2, cat_label(sr), F_BODY, border=BORDER)
    put(an, rr, 3, f"=IF($B{rr}=\"\",\"\",'{SET}'!D{sr})", F_BODY, CENTER, '0.0', border=BORDER)
    put(an, rr, 4, f'=IF($B{rr}="","",C{rr}*{C_YEAR})', F_BOLD, CENTER, '#,##0', border=BORDER)
    put(an, rr, 5, f"=IF($B{rr}=\"\",\"\",D{rr}/'{SET}'!$D$6)", F_BODY, CENTER, '0.0', border=BORDER)
    put(an, rr, 6, f'=IF($B{rr}="","",D{rr}/{C_WEEK})', F_BODY, CENTER, '0.0', border=BORDER)
    put(an, rr, 7, f'=IF($B{rr}="","",C{rr}/{C_WEEK})', F_MUTED, CENTER, '0.0%', border=BORDER)
AL = A1 + N_CORE + N_CUSTOM - 1
band(an, AL + 2, 2, 7, '  💡  PERSPECTIVE')
put(an, AL + 3, 2, f"=\"At this pace you will sleep \"&TEXT('{SET}'!D{FIRST}*{C_YEAR}/24,\"0\")&\" full days this year.\"", F_SUB)
put(an, AL + 4, 2, f"=\"You will work \"&TEXT('{SET}'!D{FIRST + 1}*{C_YEAR},\"#,##0\")&\" hours — that is \"&TEXT('{SET}'!D{FIRST + 1}*{C_YEAR}/24,\"0\")&\" days.\"", F_SUB)
put(an, AL + 5, 2, f"=\"Free, unallocated time adds up to \"&TEXT(MAX(0,{FREE})*{C_YEAR},\"#,##0\")&\" hours a year. Spend them on purpose.\"", F_SUB)
an.freeze_panes = f'A{A1}'

# ============================================================================
# 📊 DASHBOARD — reference-style
# ============================================================================
db = sheet(DASH, TEAL_DARK, [2, 13, 13, 13, 13, 3, 15, 12, 12, 12, 3, 26, 10])
# hero box (like reference "Current Week")
db.merge_cells('B2:E4')
hero = db['B2']
hero.value = '👉  Current Week:\nHow do you spend your 168 hours of the week at the moment?'
hero.font = Font(name='Arial', size=13, color=INK, bold=True)
hero.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for row in db['B2:E4']:
    for c in row:
        c.fill = FILL_BOX
put(db, 5, 2, f'="Week "&WEEKNUM(TODAY(),21)&"  ·  "&TEXT(TODAY(),"dddd, d mmmm yyyy")', F_SUB)

# summary table WEEKLY / MONTHLY / ANNUALLY  (reference top-right)
put(db, 2, 7, '', F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
for i, h in enumerate(['💬 WEEKLY', '📆 MONTHLY', '🎉 ANNUALLY']):
    put(db, 2, 8 + i, h, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
put(db, 3, 7, '⌛ HOURS USED', F_HEAD, LEFT, fill=FILL_TEAL, border=BORDER)
put(db, 4, 7, '⌛ HOURS FREE', F_HEAD, LEFT, fill=FILL_TEAL, border=BORDER)
put(db, 3, 8, f'={USED}', F_BOLD, CENTER, '0.0', border=BORDER)
put(db, 3, 9, f'={USED}*{C_MONTH}', F_BOLD, CENTER, '0.0', border=BORDER)
put(db, 3, 10, f'={USED}*{C_YEAR}', F_BOLD, CENTER, '#,##0.0', border=BORDER)
put(db, 4, 8, f'={FREE}', F_BOLD, CENTER, '0.0', border=BORDER)
put(db, 4, 9, f'={FREE}*{C_MONTH}', F_BOLD, CENTER, '0.0', border=BORDER)
put(db, 4, 10, f'={FREE}*{C_YEAR}', F_BOLD, CENTER, '#,##0.0', border=BORDER)
db.conditional_formatting.add('H4:J4', FormulaRule(
    formula=['$H$4<0'], fill=PatternFill('solid', fgColor=RED_BG),
    font=Font(name='Arial', size=10, color=RED, bold=True)))

# LIFE AREA / HOURS table (reference right panel) — cols L,M
LT = 7
put(db, LT, 12, 'LIFE AREA', F_HEAD, LEFT, fill=FILL_TEAL, border=BORDER)
put(db, LT, 13, 'HOURS', F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
for i in range(N_CORE):
    rr, sr = LT + 1 + i, FIRST + i
    put(db, rr, 12, cat_label(sr), F_BODY, border=BORDER)
    put(db, rr, 13, f"='{SET}'!D{sr}", F_BODY, CENTER, '0.0', border=BORDER)
FREE_ROW = LT + 1 + N_CORE
fc = put(db, FREE_ROW, 12, '✅  Free Time', F_BOLD, border=BORDER)
fv = put(db, FREE_ROW, 13, f'=MAX(0,{FREE})', F_BOLD, CENTER, '0.0', border=BORDER)
fc.font = fv.font = Font(name='Arial', size=10, bold=True, color=GREEN)
for i in range(N_CUSTOM):
    rr, sr = FREE_ROW + 1 + i, FIRST + N_CORE + i
    put(db, rr, 12, cat_label(sr), F_MUTED, border=BORDER)
    put(db, rr, 13, f"=IF('{SET}'!C{sr}=\"\",\"\",'{SET}'!D{sr})", F_MUTED, CENTER, '0.0', border=BORDER)
TOT_ROW = FREE_ROW + N_CUSTOM + 1
tc = put(db, TOT_ROW, 12, '❌  Total Non-Free Time', F_BOLD, border=BORDER)
tv = put(db, TOT_ROW, 13, f'={USED}', F_BOLD, CENTER, '0.0', border=BORDER)
tc.font = tv.font = Font(name='Arial', size=10, bold=True, color=RED)

# charts (data = life-area table incl. Free Time row)
data = Reference(db, min_col=13, min_row=LT, max_row=FREE_ROW)      # header + rows
cats = Reference(db, min_col=12, min_row=LT + 1, max_row=FREE_ROW)
pie = PieChart()
pie.title = 'What does my time split look like by %?'
pie.add_data(data, titles_from_data=True)
pie.set_categories(cats)
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
pie.height, pie.width = 8.6, 10.5
pie.style = 10
db.add_chart(pie, 'B7')

bar = BarChart()
bar.type = 'bar'
bar.title = 'What does my time split look like overall?'
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.legend = None
bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True; bar.dataLabels.numFmt = '0.0'
bar.height, bar.width = 8.6, 10.5
s = bar.series[0]
s.graphicalProperties.solidFill = ORANGE
db.add_chart(bar, 'G7')

# signals + score row (below charts)
SG = 26
band(db, SG, 2, 5, '  🚦  SIGNALS')
sigs = [
    f'=IF(\'{SET}\'!D{FIRST + 1}>60,"🔴  Work is over 60 h/week — burnout territory. Cut or delegate.",IF(\'{SET}\'!D{FIRST + 1}>55,"🟡  Work is above 55 h/week — watch the load.","🟢  Work load is sustainable."))',
    f'=IF(\'{SET}\'!D{FIRST}<49,"🔴  Sleep is under 7 h/night — recovery is your first lever.","🟢  Sleep is on target.")',
    f'=IF(\'{SET}\'!D{FIRST + 5}<3,"🟡  Fitness is under 3 h/week — schedule workouts first, not last.","🟢  Fitness habit is funded.")',
    f'=IF({FREE}<0,"🔴  You have allocated MORE than 168 hours — something must give.",IF({FREE}<5,"🟡  Under 5 h of true free time — build in slack.","🟢  Healthy buffer of free time."))',
    f'=IF(\'{SET}\'!D{FIRST + 7}<3,"🟡  Relationships get under 3 h/week — book the quality time.","🟢  Relationships are funded.")',
]
for i, f in enumerate(sigs):
    db.merge_cells(start_row=SG + 1 + i, start_column=2, end_row=SG + 1 + i, end_column=5)
    put(db, SG + 1 + i, 2, f, F_BODY)

band(db, SG, 7, 10, '  🔭  HORIZONS')
for i, h in enumerate(['', 'Daily', 'Weekly', 'Monthly', 'Yearly'][1:]):
    put(db, SG + 1, 7 + i if i else 8, h, F_BOLD, CENTER)
put(db, SG + 1, 7, '', F_BOLD)
hz = [('Hours used', USED), ('Free time', f'MAX(0,{FREE})'),
      ('Productive', PROD), ('Sleep', f"'{SET}'!D{FIRST}")]
for i, (lab, ref) in enumerate(hz):
    rr = SG + 2 + i
    put(db, rr, 7, lab, F_MUTED)
    put(db, rr, 8, f'={ref}/7', F_BODY, CENTER, '0.0')
    put(db, rr, 9, f'={ref}', F_BODY, CENTER, '0.0')
    put(db, rr, 10, f'={ref}*{C_MONTH}', F_BODY, CENTER, '#,##0.0')
    put(db, rr, 11, f'={ref}*{C_YEAR}', F_BODY, CENTER, '#,##0')
for row in range(SG + 1, SG + 6):
    for col in range(7, 12):
        db.cell(row=row, column=col).border = BORDER

SC = SG + 8
cards = [
    ('🎯 BALANCE SCORE',
     f'=ROUND(100*((\'{SET}\'!D{FIRST}>=49)*(\'{SET}\'!D{FIRST}<=63)'
     f'+(\'{SET}\'!D{FIRST + 1}<=55)*1+(\'{SET}\'!D{FIRST + 5}>=3)*1'
     f'+(\'{SET}\'!D{FIRST + 7}>=3)*1+((\'{SET}\'!D{FIRST + 8}+\'{SET}\'!D{FIRST + 9})>=2)*1'
     f'+({FREE}>=5)*1+({FREE}>=0)*1)/7,0)', 'out of 100'),
    ('⚡ UTILIZATION', f'={USED}/{C_WEEK}', 'of your 168 hours'),
    ('💼 PRODUCTIVE', f'={PROD}', 'productive h / week'),
]
for i, (lab, formula, sub) in enumerate(cards):
    c0 = 2 + i * 2
    put(db, SC, c0, lab, F_HEAD, CENTER, fill=FILL_TEAL, border=BORDER)
    db.merge_cells(start_row=SC, start_column=c0, end_row=SC, end_column=c0 + 1)
    v = put(db, SC + 1, c0, formula, F_KPI, CENTER,
            '0.0%' if 'UTILIZATION' in lab else '0.0' if 'PROD' in lab else '0')
    db.merge_cells(start_row=SC + 1, start_column=c0, end_row=SC + 1, end_column=c0 + 1)
    put(db, SC + 2, c0, sub, F_MUTED, CENTER)
    db.merge_cells(start_row=SC + 2, start_column=c0, end_row=SC + 2, end_column=c0 + 1)
    for rr in range(SC, SC + 3):
        for cc in range(c0, c0 + 2):
            db.cell(row=rr, column=cc).border = BORDER
            if rr > SC:
                db.cell(row=rr, column=cc).fill = FILL_CARD
BAL_CELL = f'B{SC + 1}'
db.conditional_formatting.add(BAL_CELL, FormulaRule(
    formula=[f'${BAL_CELL[0]}${BAL_CELL[1:]}>=80'],
    font=Font(name='Arial', size=20, bold=True, color=GREEN)))
db.conditional_formatting.add(BAL_CELL, FormulaRule(
    formula=[f'${BAL_CELL[0]}${BAL_CELL[1:]}<60'],
    font=Font(name='Arial', size=20, bold=True, color=RED)))

# ============================================================================
# 📖 START HERE
# ============================================================================
gd = sheet(GUI, 'B7B7B7', [2, 4, 110])
put(gd, 2, 2, '📖 Start Here — how to run your 168 hours', F_TITLE)
rows = [
    ('🧭  THE IDEA', [
        'Everyone gets exactly 24 hours a day and 168 hours a week. This workbook treats those 168 hours like a budget: you allocate them on 📝 Weekly Tracker, log reality on 📅 Daily Tracker, and everything reconciles automatically.']),
    ('✏️  RULE #1 — ONLY EDIT LIGHT-BLUE CELLS', [
        'Every editable cell has a light-blue background. Everything else is a formula — you never need to touch it.']),
    ('🚀  GETTING STARTED (10 MINUTES)', [
        '1.  Go to 📝 Weekly Tracker and answer the guided sentences (blue cells). The sheet already contains realistic example values — replace them with yours.',
        '2.  Watch 📊 Dashboard — HOURS USED / HOURS FREE, the charts and the Signals update live.',
        '3.  During the week, log real hours on 📅 Daily Tracker.',
        '4.  Review 📈 Weekly Review at week end: planned vs actual, and where the gap is.']),
    ('➕  ADDING A CUSTOM CATEGORY', [
        'Open ⚙️ Settings, find the light-blue rows at the bottom of the Master Category Table, and type an emoji, a name and weekly hours. It appears in every tracker and the Dashboard automatically.']),
    ('🎯  BALANCE SCORE', [
        'A 0–100 gauge of how humane the week is: sleep 49–63 h, work ≤ 55 h, fitness ≥ 3 h, relationships ≥ 3 h, learning + personal development ≥ 2 h, free time ≥ 5 h, nothing over-allocated. 80+ = sustainable.']),
    ('🔁  A NEW WEEK', [
        'Select the light-blue grid on 📅 Daily Tracker and press Delete. Your plan stays.']),
    ('📐  ASSUMPTIONS', [
        'Monthly = weekly × 4, yearly = weekly × 52 (the reference sheet\'s convention). Both constants are editable on ⚙️ Settings — use 4.345 / 52.18 for calendar-accurate maths. Default input values reproduce the reference screenshots (166.1 h used, 1.9 h free).']),
]
gr = 4
for title, paras in rows:
    band(gd, gr, 2, 3, '  ' + title)
    gr += 1
    for p in paras:
        cell = put(gd, gr, 3, p, F_MUTED)
        cell.alignment = WRAP
        gd.row_dimensions[gr].height = max(16, 15 * (1 + len(p) // 105))
        gr += 1
    gr += 1

# tab order: Dashboard first
wb.move_sheet(DASH, offset=-(len(wb.sheetnames) - 1))
wb.active = wb[DASH]
wb.save(OUT)
print('saved', OUT)

# quick self-check of key wiring
import subprocess, json
res = subprocess.run(['python3', '/root/.claude/skills/xlsx/scripts/recalc.py', OUT, '60'],
                     capture_output=True, text=True)
print(res.stdout[-2000:] or res.stderr[-2000:])
